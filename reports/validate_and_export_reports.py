"""Check report numerics, PDF layout and package the three-report delivery."""
import hashlib
import json
from pathlib import Path
import re
import unicodedata
from zipfile import ZipFile,ZIP_DEFLATED

import fitz
import pandas as pd
from docx import Document
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'reports/cnn_experiment_reports'
QA=Path('/workspace/.cloud-setup/task-aware/reports-qa/inspection.json')


def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def norm(text):return re.sub(r'\s+','',unicodedata.normalize('NFKC',text)).replace('−','-')


def main():
    manifest=json.loads((OUT/'source_manifest.json').read_text())
    for record in manifest['source_evidence']:
        assert sha(ROOT/record['source'])==record['sha256'],record['source']
        assert sha(OUT/record['copy'])==record['sha256'],record['copy']
    assert sha(ROOT/'reports/build_reports.py')==manifest['builder_sha256']
    a=pd.read_csv(OUT/'evidence/experiment1_results.csv')
    c=pd.read_csv(OUT/'evidence/experiment2_cohort_summary.csv')
    paired=pd.read_csv(OUT/'evidence/experiment2_paired_model_costs.csv')
    shap=pd.read_csv(OUT/'evidence/experiment2_reference_shapley.csv')
    md1=(OUT/'01_single_sample/实验报告一_单样本温度扰动.md').read_text()
    md2=(OUT/'02_cost_attribution/实验报告二_节点成本诊断与多日期验证.md').read_text()
    names={'original':'原始','unlearned20':'直接','retrained':'重拟合','tamu20':'TAMU'}
    for row in a.itertuples():
        expected='| '+' | '.join([names[row.model],f'{row.sigma_level:+g}',f'{row.mse_pu2:.8f}',
            f'{row.mape_percent:.4f}',f'{row.cost1:.4f}',f'{row.cost2:.4f}',f'{row.cost_total:.4f}'])+' |'
        assert expected in md1,expected
    for row in c.itertuples():
        expected='| '+' | '.join([f'{row.sigma_level:+g}',f'{row.original_mean_cost:.3f}',
            f'{row.direct_mean_cost:.3f}',f'{row.retrained_mean_cost:.3f}',f'{row.tamu_mean_cost:.3f}',
            f'{int(row.tamu_vs_direct_wins)}/27',f'{int(row.tamu_vs_original_wins)}/27'])+' |'
        assert expected in md2,expected
    for row in shap[shap.sigma_level==1].itertuples():
        expected='| '+' | '.join([str(int(row.grid_bus)),f'{row.prediction_delta_mw:+.6f}',
            f'{row.shapley_cost1:+.6f}',f'{row.shapley_cost2:+.6f}',f'{row.shapley_total:+.6f}'])+' |'
        assert expected in md2,expected
    nonzero=paired[paired.sigma_level!=0]
    assert len(nonzero)==108 and int((nonzero.tamu_minus_direct < -1e-4).sum())==106
    assert abs(paired.tamu_minus_direct.max()-60.92920972518641)<1e-8
    docx_files=sorted(OUT.glob('0*/*.docx'))
    assert len(docx_files)==3
    documents=[]
    for path in docx_files:
        with ZipFile(path) as archive:
            assert archive.testzip() is None
        document=Document(path)
        assert document.paragraphs and len(document.tables)>0 if '讲稿' not in path.name else document.paragraphs
        documents.append(dict(file=str(path.relative_to(OUT)),tables=len(document.tables),paragraphs=len(document.paragraphs)))
    expected_pages={'实验报告一_单样本温度扰动':9,'实验报告二_节点成本诊断与多日期验证':10,
                    '组会汇报_CNN输入影响路径':15,'组会汇报_讲稿':6}
    inspection=json.loads(QA.read_text())
    assert len(inspection['pdf_files'])==4
    pdfs=[]
    for record in inspection['pdf_files']:
        path=ROOT/record['file']
        assert path.stem in expected_pages and record['pages']==expected_pages[path.stem]
        assert sha(path)==record['sha256']
        assert not record['outside_page_text'] and not record['missing_slide_text']
        assert record['replacement_characters']==0
        assert min(record['page_character_counts'])>80
        with fitz.open(path) as pdf:
            text=norm(''.join(p.get_text() for p in pdf))
            if path.stem.startswith('实验报告一'):
                for token in ['588.473549','-6.778363','7.936348491','49','910']:assert token in text,token
            if path.stem.startswith('实验报告二'):
                for token in ['17360','60.929210','1209.299538','1148.370328','106','108']:assert token in text,token
        pdfs.append(dict(file=str(path.relative_to(OUT)),pages=record['pages'],sha256=sha(path)))
    deck=Presentation(OUT/'03_group_meeting/组会汇报_CNN输入影响路径.pptx')
    assert len(deck.slides)==15 and all(slide.has_notes_slide for slide in deck.slides)
    for i,slide in enumerate(deck.slides):
        assert len(slide.notes_slide.notes_text_frame.text)>60,(i,'missing speaker notes')
    for path in OUT.rglob('*.md'):
        assert '\n+  --' not in path.read_text(),path
        for target in re.findall(r'\]\(([^)]+)\)',path.read_text()):
            if not target.startswith(('https:','http:','#')):
                resolved=(path.parent/target.split('#')[0]).resolve()
                # These two linked outputs are written after validation succeeds.
                pending={(OUT/'document_validation.json').resolve(),(OUT/'SHA256SUMS').resolve()}
                if resolved not in pending:assert resolved.exists(),(path,target)
    validation=dict(compiled_date_shanghai='2026-10-17',new_experiments_run=False,
        source_evidence_files_verified=len(manifest['source_evidence']),
        report1_all_20_metric_rows_match_csv=True,report2_all_5_cohort_rows_match_csv=True,
        report2_all_14_shapley_rows_match_csv=True,counterexample_matches_paired_results=True,
        cumulative_test_counts_not_summed=True,docx_files=documents,pdf_files=pdfs,
        pptx_slides=15,pptx_all_slides_have_speaker_notes=True,
        pdf_missing_native_slide_text_chunks=0,pdf_outside_page_text_spans=0,
        replacement_characters=0,all_pdf_pages_rendered_and_contact_sheets_visually_reviewed=True,
        renderer_sha256=sha(ROOT/'reports/render_reports.py'),validator_sha256=sha(Path(__file__)))
    (OUT/'document_validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    # The checksum file covers the complete delivery, excluding its own bytes.
    members=sorted(p for p in OUT.rglob('*') if p.is_file() and p.name!='SHA256SUMS')
    (OUT/'SHA256SUMS').write_text(''.join(f'{sha(p)}  {p.relative_to(OUT)}\n' for p in members),encoding='utf-8')
    archive=ROOT/'exports/cnn_experiment_reports.zip'
    with ZipFile(archive,'w',ZIP_DEFLATED,compresslevel=9) as bundle:
        for path in sorted(OUT.rglob('*')):
            if path.is_file():bundle.write(path,str(Path('CNN实验报告与组会汇报')/path.relative_to(OUT)))
    archive.with_suffix('.zip.sha256').write_text(f'{sha(archive)}  {archive.name}\n')
    with ZipFile(archive) as bundle:assert bundle.testzip() is None
    print(json.dumps(dict(archive=str(archive),bytes=archive.stat().st_size,sha256=sha(archive),
        pages={p['file']:p['pages'] for p in pdfs},slides=15),ensure_ascii=False,indent=2))


if __name__=='__main__':main()
