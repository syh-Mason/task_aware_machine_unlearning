"""Render editable documents with LibreOffice and inspect their PDF outputs."""
import argparse
import hashlib
import json
import os
from pathlib import Path
import re
import subprocess
import unicodedata
from zipfile import ZipFile

import fitz
from PIL import Image,ImageOps,ImageDraw
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'reports/cnn_experiment_reports'
QA=Path('/workspace/.cloud-setup/task-aware/reports-qa')


def norm(text):
    return re.sub(r'\s+','',unicodedata.normalize('NFKC',text)).replace('−','-').replace('–','-')


def render():
    files=sorted(OUT.glob('*/*.docx'))+sorted(OUT.glob('*/*.pptx'))
    assert len(files)==4
    for index,path in enumerate(files):
        profile=Path(f'/tmp/cnn-reports-lo-{index}');profile.mkdir(exist_ok=True)
        command=['soffice',f'-env:UserInstallation={profile.as_uri()}',
            '--headless','--convert-to','pdf','--outdir',str(path.parent),str(path)]
        result=subprocess.run(command,capture_output=True,text=True,env={**os.environ,'XDG_CACHE_HOME':'/workspace/.cache'},timeout=90)
        print(result.stdout.strip(),flush=True)
        if result.returncode or not path.with_suffix('.pdf').exists():
            raise RuntimeError(result.stdout+'\n'+result.stderr)


def inspect():
    QA.mkdir(parents=True,exist_ok=True)
    summaries=[];all_contacts=[]
    for document in sorted(OUT.glob('0*/*.pdf')):
        with fitz.open(document) as pdf:
            text=[page.get_text() for page in pdf]
            previews=[];outside=[]
            for number,page in enumerate(pdf):
                image=page.get_pixmap(matrix=fitz.Matrix(.8,.8),alpha=False)
                target=QA/f'{document.parent.name}-{document.stem}-{number+1:02d}.png'
                image.save(target);previews.append(target)
                for block in page.get_text('dict')['blocks']:
                    if block['type']!=0:continue
                    for line in block['lines']:
                        for span in line['spans']:
                            rect=fitz.Rect(span['bbox'])
                            if rect.x0<-.8 or rect.y0<-.8 or rect.x1>page.rect.width+.8 or rect.y1>page.rect.height+.8:
                                outside.append(dict(page=number+1,text=span['text'],bbox=list(rect)))
            if document.stem.startswith('组会汇报_CNN'):
                prs=Presentation(document.with_suffix('.pptx'))
                assert len(pdf)==len(prs.slides)==15
                missing=[]
                for i,slide in enumerate(prs.slides):
                    available=norm(text[i])
                    for shape in slide.shapes:
                        chunks=[]
                        if shape.has_text_frame:chunks.append(shape.text)
                        if shape.has_table:
                            chunks.extend(cell.text for row in shape.table.rows for cell in row.cells)
                        for chunk in chunks:
                            if norm(chunk) and norm(chunk) not in available:
                                missing.append(dict(slide=i+1,text=chunk))
                (QA/'slide_text_missing.json').write_text(json.dumps(missing,ensure_ascii=False,indent=2))
            else:missing=[]
            record=dict(file=str(document.relative_to(ROOT)),pages=len(pdf),
                page_character_counts=[len(t.strip()) for t in text],
                outside_page_text=outside,missing_slide_text=missing,
                replacement_characters=sum(t.count('\ufffd') for t in text),
                sha256=hashlib.sha256(document.read_bytes()).hexdigest())
            summaries.append(record)
            # Contact sheets are QA previews, not experimental figures.
            thumb_w=320;thumb_h=440 if pdf[0].rect.height>pdf[0].rect.width else 205
            cols=3;rows=(len(previews)+cols-1)//cols
            sheet=Image.new('RGB',(cols*thumb_w,rows*(thumb_h+28)),(220,225,229));draw=ImageDraw.Draw(sheet)
            for i,path in enumerate(previews):
                with Image.open(path) as im:
                    thumb=ImageOps.contain(im,(thumb_w-10,thumb_h-4))
                    x=(i%cols)*thumb_w+(thumb_w-thumb.width)//2;y=(i//cols)*(thumb_h+28)
                    sheet.paste(thumb,(x,y));draw.text(((i%cols)*thumb_w+8,y+thumb_h+3),f'Page {i+1}',fill=(0,0,0))
            target=QA/f'{document.parent.name}-{document.stem}-contact.png';sheet.save(target)
            all_contacts.append(str(target))
    result=dict(pdf_files=summaries,contact_sheets=all_contacts)
    (QA/'inspection.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps([dict(file=x['file'],pages=x['pages'],characters=x['page_character_counts'],
        outside=len(x['outside_page_text']),missing_slide_chunks=len(x['missing_slide_text'])) for x in summaries],ensure_ascii=False,indent=2))


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--inspect-only',action='store_true')
    args=parser.parse_args()
    if not args.inspect_only:render()
    inspect()


if __name__=='__main__':main()
