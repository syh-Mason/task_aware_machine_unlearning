"""Visualize measured CNN module paths and internal attributions."""
import csv
import hashlib
import json
import os
from pathlib import Path
import subprocess
from zipfile import ZipFile, ZIP_DEFLATED

os.environ.setdefault('XDG_CACHE_HOME','/workspace/.cache')
os.environ.setdefault('MPLCONFIGDIR','/workspace/.cloud-setup/task-aware/matplotlib')
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
import pandas as pd
from PIL import Image,ImageDraw,ImageOps
import fitz
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'research/cnn_layer_paths';FIG=OUT/'figures'
QA=Path('/workspace/.cloud-setup/task-aware/cnn-layer-qa')
NAMES={'original':'原始CNN','unlearned20':'直接遗忘20%','retrained':'剩余集重拟合头','tamu20':'TAMU'}
FONT='Noto Sans CJK SC'


def plots():
    fontpath='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
    font_manager.fontManager.addfont(fontpath)
    name=font_manager.FontProperties(fname=fontpath).get_name()
    plt.rcParams.update({'font.family':name,'font.size':12,'axes.unicode_minus':False,'pdf.fonttype':42})
    a=np.load(OUT/'layer_arrays.npz')
    def save(fig,name):
        for ext in ['png','pdf']:fig.savefig(FIG/(name+'.'+ext),dpi=200,bbox_inches='tight',facecolor='white')
        plt.close(fig)
    for layer,title in [('relu1','Conv1 + ReLU：8个通道激活'),('relu2','Conv2 + ReLU：8个通道激活')]:
        values=a['original.activation.'+layer][0]
        fig,axes=plt.subplots(2,4,figsize=(10,5),layout='constrained')
        for i,ax in enumerate(axes.flat):
            im=ax.imshow(values[i],aspect='auto',cmap='viridis',vmin=0,vmax=values.max())
            ax.set_title('通道 '+str(i+1));ax.set_xticks([]);ax.set_yticks([])
        fig.colorbar(im,ax=axes.ravel().tolist(),label='激活值')
        fig.suptitle(title+'（四模型共享同一激活）')
        save(fig,'02_'+layer+'_activation')
    # Scores for the same hidden layer are comparable in the same output units.
    fig,axes=plt.subplots(2,2,figsize=(12,7.5),layout='constrained')
    bound=max(np.abs(a[m+'.head_contribution_mw']).max() for m in NAMES)
    for ax,(m,label) in zip(axes.flat,NAMES.items()):
        im=ax.imshow(a[m+'.head_contribution_mw'],aspect='auto',cmap='RdBu_r',vmin=-bound,vmax=bound)
        ax.set(title=label,xlabel='隐藏神经元编号（1—64）',ylabel='输出节点',yticks=[0,3,7,13],yticklabels=[1,4,8,14],xticks=[0,15,31,47,63],xticklabels=[1,16,32,48,64])
    fig.colorbar(im,ax=axes.ravel().tolist(),label='W[j,k] × h[k] × std[j] × 100 / MW')
    save(fig,'03_hidden_to_output_contributions')
    fig,axes=plt.subplots(2,2,figsize=(11,7),layout='constrained')
    for ax,(m,label) in zip(axes.flat,NAMES.items()):
        value=a[m+'.head_contribution_mw'][3];order=np.argsort(-np.abs(value))[:10]
        ax.bar(range(10),value[order],color=['#B2533B' if x<0 else '#2460A6' for x in value[order]])
        ax.set(title=label+'：节点4',xlabel='神经元编号（按绝对贡献排序）',ylabel='带符号贡献 / MW',xticks=range(10),xticklabels=order+1)
        ax.axhline(0,color='0.4',lw=.6);ax.grid(axis='y',alpha=.2)
    save(fig,'04_node4_top_neurons')
    rows=pd.read_csv(OUT/'internal_ablations.csv')
    fig,axes=plt.subplots(1,3,figsize=(13,4.2),layout='constrained',gridspec_kw={'width_ratios':[1,1,3]})
    for ax,layer,title in zip(axes,['relu1','relu2','tanh'],['ReLU1：整通道置零','ReLU2：整通道置零','Tanh：单隐藏神经元置零']):
        table=rows[(rows.model=='original')&(rows.layer==layer)].pivot(index='output_bus',columns='unit_label1',values='delta_prediction_mw')
        bound=np.abs(table.values).max()
        im=ax.imshow(table,aspect='auto',cmap='RdBu_r',vmin=-bound,vmax=bound)
        ax.set(title=title,xlabel='通道/神经元编号',ylabel='输出节点',yticks=[0,3,7,13],yticklabels=[1,4,8,14])
        fig.colorbar(im,ax=ax,label='置零后预测 − 原预测 / MW')
    save(fig,'05_internal_ablation_effects')


def diagram():
    labels=[('input','固定未扰动样本\n(N,1,14,10)'),
        ('conv1','Conv1｜3×3 / stride 1\n(N,8,14,10)｜80参数\n局部特征提取'),
        ('relu1','ReLU 第1次调用\n(N,8,14,10)\n本样本48.66%激活为零'),
        ('conv2','Conv2｜4×4 / stride 2\n(N,8,7,5)｜1032参数\n局部感受野扩为6×6编码位置'),
        ('relu2','ReLU 第2次调用\n(N,8,7,5)\n本样本33.93%激活为零'),
        ('flatten','Flatten\n(N,8,7,5) → (N,280)\n只重排，不学习新特征'),
        ('lin1','全连接 lin1｜280→64\n(N,64)｜17984参数\n融合所有卷积位置'),
        ('tanh','Tanh 隐藏表示 h\n(N,64)｜限制到(-1,1)\n本样本没有 |h|>0.95 的单元'),
        ('head','仿射预测头｜64→14\n(N,14)｜910参数\n原始/直接遗忘/重拟合/TAMU 各有一组参数'),
        ('output','14节点预测\n反标准化 × baseMVA → MW')]
    lines=['digraph layers {','graph [rankdir=TB,bgcolor="white",nodesep=.2,ranksep=.24,pad=.2];',
        'node [shape=box,style="rounded,filled",fillcolor="#EFF5F8",color="#167D8D",fontname="Noto Sans CJK SC",fontsize=18,margin=".18,.10"];',
        'edge [color="#526675",penwidth=1.4];']
    for key,label in labels:lines.append(key+' [label='+json.dumps(label,ensure_ascii=False)+'];')
    lines.append(' -> '.join(k for k,_ in labels)+';')
    lines.append('subgraph cluster_core {label="四模型共享且冻结的核心：内部激活相同";fontname="Noto Sans CJK SC";fontsize=16;color="#167D8D";style=dashed;conv1;relu1;conv2;relu2;flatten;lin1;tanh;}')
    lines.append('head [fillcolor="#F5E9E5",color="#B2533B"];')
    lines.append('graph [label="模型内部模块路径｜Captum层归因 + 内部通道/神经元置零｜固定单样本；不解释为物理因果",labelloc=b,fontname="Noto Sans CJK SC",fontsize=11];}')
    (FIG/'01_model_layers.dot').write_text('\n'.join(lines)+'\n')
    for ext in ['svg','pdf','png']:
        command=['dot','-T'+ext,str(FIG/'01_model_layers.dot'),'-o',str(FIG/('01_model_layers.'+ext))]
        if ext=='png':command.insert(1,'-Gdpi=150')
        subprocess.run(command,check=True,capture_output=True)


def report():
    validation=json.loads((OUT/'validation.json').read_text())
    stats=pd.read_csv(OUT/'layer_statistics.csv');a=np.load(OUT/'layer_arrays.npz')
    doc=Document();s=doc.sections[0]
    s.page_width=Inches(8.2677);s.page_height=Inches(11.6929)
    s.left_margin=s.right_margin=Inches(.7);s.top_margin=s.bottom_margin=Inches(.65)
    for name,size in [('Normal',10.5),('Title',23),('Heading 1',16),('Heading 2',12),('Caption',9)]:
        st=doc.styles[name];st.font.name=FONT;st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT);st.font.size=Pt(size)
    doc.styles['Normal'].paragraph_format.line_spacing=1.18
    s.header.paragraphs[0].text='CNN 模型内部层作用路径｜固定样本分析'
    p=s.footer.paragraphs[0];p.alignment=1
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
    md=['# CNN 模型内部层作用路径分析','',
        '[下载Word](CNN_模型内部层作用路径.docx?raw=true) · [阅读PDF](CNN_模型内部层作用路径.pdf) · [下载完整包](../../exports/cnn_layer_paths.zip?raw=true) · [复算说明](REPRODUCE.md)','']
    def page(title,first=False):
        if not first:doc.add_page_break()
        doc.add_heading(title,1);md.extend(['## '+title,''])
    def p(text):doc.add_paragraph(text);md.extend([text,''])
    def image(name,caption,maxheight=4.8):
        path=FIG/(name+'.png')
        with Image.open(path) as im:w,h=im.size;scale=min(6.8/w,maxheight/h)
        doc.add_picture(str(path),width=Inches(w*scale));doc.add_paragraph(caption,style='Caption')
        md.extend(['!['+caption+'](figures/'+name+'.png)',''])
    def table(headers,rows):
        t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
        for c,text in zip(t.rows[0].cells,headers):c.text=str(text)
        for row in rows:
            for c,text in zip(t.add_row().cells,row):c.text=str(text)
        for r in t.rows:
            for c in r.cells:
                for par in c.paragraphs:
                    for run in par.runs:run.font.size=Pt(9)
        md.extend(['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |'])
        md.extend('| '+' | '.join(map(str,row))+' |' for row in rows);md.append('')
    page('CNN 各层怎样形成最终预测',True)
    p('分析对象：原 Task-Aware 项目的已有 CNN 与四种敏感集预测头。分析重点是模型内部各层的表示、梯度贡献和内部消融，不是再次改变温度等原始输入。日期：2026-10-21。')
    p('Captum 是归因分析工具库，本身没有唯一神经网络结构。用户上传的 captum-master.zip 超过当前工具32 MiB读取上限，未能读取；本轮使用官方 Captum v0.7.0源码，与既有 PyTorch 2.0.1兼容，不声称等同上传包版本。')
    table(['分析项','实际执行'],[['固定样本','既有 test index 0，原始未扰动输入'],['模型','原始、20%直接遗忘、剩余集重拟合头、TAMU'],['层记录','8个实际运算阶段；两次ReLU分别记录'],['层归因','Captum LayerActivation、LayerGradientXActivation'],['内部干预','ReLU1的8通道、ReLU2的8通道、Tanh的64神经元逐个置零'],['规模','4模型×80次=320次内部消融，记录4480个输出变化'],['验收',f'{validation["checks_passed"]}项内部一致性检查通过，0失败']])
    p('核心发现：四模型的核心激活逐位相同，贡献会因预测头参数改变而不同；原始模型节点4的最大绝对隐藏贡献来自神经元9，直接遗忘/重拟合/TAMU则来自神经元26。该观察仅适用于当前固定样本。')
    page('1  模型内部的实际调用链')
    image('01_model_layers','图1  冻结核心与可更新预测头；图中N为批大小，本次N=1。',7.6)
    p('原代码同一个relu模块调用两次；分析包装器将两次调用分成唯一命名relu1/relu2，保持计算与预测逐位一致，避免层hook混淆。')
    page('2  每层做什么、形状怎样变化')
    table(['阶段','实际输出形状','参数数'],[[r.layer,r['shape'],r.parameters] for _,r in stats[stats.model=='original'].iterrows()])
    p('Conv1在编码张量的局部3×3邻域提取8种响应；Conv2组合已有通道并以步长2缩小空间尺寸，局部感受野约6×6编码位置。这里的相邻位置由bus/特征排列决定，不能直接称为电网物理相邻。')
    p('ReLU只保留正响应。Flatten仅将8×7×5重排为280维；lin1把所有位置融合到64维；Tanh将表示限制到(-1,1)，可能因饱和而减弱梯度；当前样本没有绝对值大于0.95的Tanh单元。')
    p('最终预测头以Wx+b将64维隐藏表示读为14个负荷输出，再反标准化为p.u.并换算MW。实际有效网络共20006参数，其中预测头910参数。本轮没有重新训练这些参数。')
    page('3  看见各卷积通道的激活')
    image('02_relu1_activation','图2  第一层卷积经ReLU后的8个通道；行列对应编码位置。',3.3)
    image('02_relu2_activation','图3  第二层卷积经ReLU后的8个通道。',3.3)
    p('当前样本ReLU1零激活比例48.6607%，ReLU2为33.9286%。这描述哪些内部响应被门控，不表示通道已经学到某个已知物理含义。四模型核心相同，所以这些激活也相同。')
    page('4  哪些隐藏神经元在支持各输出')
    image('03_hidden_to_output_contributions','图4  四模型的64隐藏神经元→14输出贡献，单位MW。',4.7)
    p('Captum的LayerGradientXActivation计算“该层激活×目标输出对该层激活的梯度”。这里目标是每个节点的原始负荷预测（MW），不是Cost。不同层的分数是同一预测的不同切面，不能相加或据其总量给层做全局重要性排名。')
    p('对最后仿射头，隐藏贡献具有精确表达：贡献[j,k]=W[j,k]×h[k]×std[j]×100。所有隐藏贡献加上(b[j]×std[j]+mean[j])×100，恢复节点预测。当前最大重构误差为1.43e−14 MW。')
    p('这是对隐藏状态置零参考下的仿射读出分解；常数偏置单列。前面非线性层的Gradient×Activation一般只是局部归因，不自动具备同样的精确可加性。')
    page('5  节点4：遗忘前后读出了什么')
    image('04_node4_top_neurons','图5  节点4的前10个绝对隐藏贡献；符号表示抬高或拉低该预测。',4.5)
    rows=[]
    for m,name in NAMES.items():
        c=a[m+'.head_contribution_mw'][3];i=int(np.argmax(np.abs(c)))
        rows.append([name,f'{a[m+".prediction_mw"][0,3]:.6f}',i+1,f'{c[i]:+.6f}'])
    table(['模型','节点4预测/MW','最大绝对贡献神经元','带符号贡献/MW'],rows)
    p('原始模型神经元9贡献−5.526458 MW；直接遗忘与重拟合的最大贡献转为神经元26，约−8.8277 MW；TAMU同为神经元26，约−6.576171 MW。前层激活没变，改变的是预测头对同一隐藏表示的读出。不能据此称核心学会了新的特征，也不能把神经元9直接命名为温度神经元。')
    page('6  内部消融验证与结论范围')
    image('05_internal_ablation_effects','图6  原始模型内部通道/神经元置零后，14个预测怎样变化；各面板色标独立。',3.3)
    p('保持原始输入不变，将某个ReLU后的整通道或某个Tanh神经元激活置0，再计算后续层。整通道包含许多空间位置，与单个隐藏神经元不是同样大小的干预，不能直接将两类峰值当作公平层排名。')
    p('原始模型将神经元9置零，节点4预测从51.640998增加到57.167456 MW，增加5.526458 MW，与其负贡献被移除一致。全部隐藏神经元置零检查的最大误差约3.27e−14 MW。前层通道消融可能改变后续ReLU门控，未必等于局部Gradient×Activation。')
    p('Captum激活与手动记录完全一致；Captum梯度×激活与直接自动微分完全一致；四模型预测与既有零扰动记录一致；模型参数与源快照均未变化。90项检查为本轮内部一致性验收，不与此前实验测试数累加。')
    p('局限：仅一条固定样本；置零后的隐藏状态未必来自真实输入；归因依赖目标和方法；没有枚举完整的最重要神经元连通路径，没有宣称物理因果或全局层重要性。本轮不重新计算Cost，也没有把数据输入扰动冒充模型层分析。')
    p('复现：使用Captum v0.7.0，运行 experiments/analyze_cnn_layers.py；源码来源 https://github.com/pytorch/captum/tree/v0.7.0 。逐层数组、归因表、消融表及校验记录保存在本目录。')
    doc.core_properties.title='CNN 模型内部层作用路径分析'
    doc.core_properties.author='task_aware_machine_unlearning 项目研究记录'
    doc.save(OUT/'CNN_模型内部层作用路径.docx')
    (OUT/'README.md').write_text('\n'.join(md))
    command=['soffice','-env:UserInstallation=file:///tmp/cnn-layer-paths-lo','--headless','--convert-to','pdf','--outdir',str(OUT),str(OUT/'CNN_模型内部层作用路径.docx')]
    r=subprocess.run(command,capture_output=True,text=True,timeout=100)
    if r.returncode:raise RuntimeError(r.stdout+'\n'+r.stderr)


def inspect_export():
    pages=[];outside=[]
    with fitz.open(OUT/'CNN_模型内部层作用路径.pdf') as pdf:
        for i,page in enumerate(pdf):
            text=page.get_text();assert len(text.strip())>40 and '\ufffd' not in text
            pages.append(len(text));page.get_pixmap(matrix=fitz.Matrix(.8,.8),alpha=False).save(QA/f'page-{i+1:02d}.png')
            for block in page.get_text('dict')['blocks']:
                if block['type']!=0:continue
                for line in block['lines']:
                    for span in line['spans']:
                        box=fitz.Rect(span['bbox'])
                        if box.x0<-.8 or box.y0<-.8 or box.x1>page.rect.width+.8 or box.y1>page.rect.height+.8:outside.append(dict(page=i+1,text=span['text']))
        print('document_pages',len(pdf),'characters',pages)
    assert not outside,outside
    (OUT/'document_validation.json').write_text(json.dumps(dict(pages=len(pages),page_character_counts=pages,outside_page_text=outside,uploaded_archive_read=False),ensure_ascii=False,indent=2)+'\n')
    thumbs=sorted(QA.glob('page-*.png'));cols=3;w=300;h=425
    contact=Image.new('RGB',(cols*w,((len(thumbs)+cols-1)//cols)*(h+24)),'#DEE5EB');draw=ImageDraw.Draw(contact)
    for i,path in enumerate(thumbs):
        with Image.open(path) as im:
            thumb=ImageOps.contain(im,(w-8,h-4));x=(i%cols)*w+(w-thumb.width)//2;y=(i//cols)*(h+24);contact.paste(thumb,(x,y));draw.text(((i%cols)*w+6,y+h),'Page '+str(i+1),fill='black')
    contact.save(QA/'contact.png')
    files=[p for p in sorted(OUT.rglob('*')) if p.is_file() and p.name!='SHA256SUMS']
    (OUT/'SHA256SUMS').write_text('\n'.join(hashlib.sha256(p.read_bytes()).hexdigest()+'  '+str(p.relative_to(OUT)) for p in files)+'\n')
    archive=ROOT/'exports/cnn_layer_paths.zip'
    with ZipFile(archive,'w',ZIP_DEFLATED) as z:
        for p in files+[OUT/'SHA256SUMS']:z.write(p,Path('CNN模型内部层分析')/p.relative_to(OUT))
        z.write(ROOT/'experiments/analyze_cnn_layers.py','CNN模型内部层分析/analyze_cnn_layers.py')
    with ZipFile(archive) as z:assert z.testzip() is None
    archive.with_suffix('.zip.sha256').write_text(hashlib.sha256(archive.read_bytes()).hexdigest()+'  '+archive.name+'\n')
    print('zip_bytes',archive.stat().st_size)


if __name__=='__main__':
    FIG.mkdir(exist_ok=True);QA.mkdir(parents=True,exist_ok=True)
    plots();diagram();report();inspect_export()
