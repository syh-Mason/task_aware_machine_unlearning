"""Build a source-bounded, editable Chinese TAMU study guide and download ZIP."""
from pathlib import Path
from copy import deepcopy
import hashlib
import json
import os
import re
import subprocess
import sys
from zipfile import ZipFile, ZIP_DEFLATED

os.environ.setdefault('XDG_CACHE_HOME', '/workspace/.cache')
os.environ.setdefault('MPLCONFIGDIR', '/workspace/.cloud-setup/task-aware/matplotlib')

import fitz
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.patches import FancyBboxPatch
import numpy as np
from PIL import Image, ImageDraw, ImageOps
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'reports/tamu_algorithm_guide'
QA = Path('/workspace/.cloud-setup/task-aware/tamu-guide-qa')
STEM = 'TAMU算法精读_先修论文与逐步推导'
FONT = 'Noto Sans CJK SC'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def assets():
    folder = OUT/'assets'
    folder.mkdir(exist_ok=True)
    font_path='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
    font_manager.fontManager.addfont(font_path)
    plot_font=font_manager.FontProperties(fname=font_path).get_name()
    plt.rcParams.update({'font.family': plot_font, 'axes.unicode_minus': False,
                         'pdf.fonttype': 42, 'font.size': 13})

    def box(ax, x, y, w, h, title, text):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle='round,pad=.008',
                                   facecolor='#EDF5F6', edgecolor='#167D8D', lw=1.3))
        ax.text(x+w/2, y+h*.73, title, fontsize=15, fontweight='bold',
                color='#19354B', ha='center', va='center')
        ax.text(x+w/2, y+h*.30, text, fontsize=12, ha='center', va='center',
                linespacing=1.6, color='#19354B')

    fig, ax = plt.subplots(figsize=(11.5, 3.8), layout='constrained')
    ax.set(xlim=(0, 1), ylim=(0, 1)); ax.axis('off')
    box(ax, .02, .26, .28, .57, 'A  直接遗忘', '剩余集 MSE\n梯度 + Hessian → Newton 更新')
    box(ax, .36, .26, .28, .57, 'B  PAMU', '影响函数评分\nMSE / MAPE → 剩余样本权重')
    box(ax, .70, .26, .28, .57, 'C  TAMU', '可微两阶段电网优化\n运行 Cost → 剩余样本权重')
    for start, end in [(.307, .35), (.647, .69)]:
        ax.annotate('', xy=(end, .55), xytext=(start, .55),
                    arrowprops=dict(arrowstyle='->', color='#19354B', lw=1.7))
    ax.text(.5, .08, '先修：影响函数 P1  →  KKT / OptNet P2  →  CVXPYLayers P3  →  任务损失 P4',
            ha='center', fontsize=13, color='#19354B')
    fig.savefig(folder/'algorithm_map.png', dpi=220, bbox_inches='tight'); plt.close(fig)

    fig, ax = plt.subplots(figsize=(11.5, 4.0), layout='constrained')
    ax.set(xlim=(0, 1), ylim=(0, 1)); ax.axis('off')
    items=[('输入 x', '14节点 × 10特征'), ('冻结 CNN + 头', '64维隐表示 → 14维预测'),
           ('阶段一调度', '预测 → 发电 + 削减'), ('阶段二再调度', '发电 + 真实负荷'),
           ('总 Cost', 'C1 + C2')]
    for i, (title, text) in enumerate(items):
        x = .01+i*.2
        box(ax, x, .44, .177, .39, title, text)
        if i < 4:
            ax.annotate('', xy=(x+.195, .64), xytext=(x+.18, .64),
                        arrowprops=dict(arrowstyle='->', color='#19354B', lw=1.5))
    ax.annotate('', xy=(.23, .31), xytext=(.90, .31),
                arrowprops=dict(arrowstyle='->', color='#B53D3D', lw=2))
    ax.text(.57, .22, 'Cost 反向梯度 → 参数梯度 v → iHVP → 样本分数 → 权重 LP',
            color='#B53D3D', ha='center', fontsize=13)
    ax.text(.5, .05, '温度扰动实验改变输入 x；TAMU 改变训练样本权重 w。二者通过同一预测—优化路径连接。',
            ha='center', fontsize=12, color='#19354B')
    fig.savefig(folder/'forward_backward_path.png', dpi=220, bbox_inches='tight'); plt.close(fig)


def reference():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.2677); section.page_height = Inches(11.6929)
    section.left_margin = section.right_margin = Inches(.76)
    section.top_margin = Inches(.72); section.bottom_margin = Inches(.68)
    section.header_distance = section.footer_distance = Inches(.30)
    for style in doc.styles:
        if style.type == 1 or style.type == 2:
            style.font.name = FONT
            style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    for name in ['Normal', 'Body Text', 'First Paragraph']:
        if name not in doc.styles:
            continue
        s = doc.styles[name]; s.font.size = Pt(10.5)
        s.paragraph_format.line_spacing = 1.23
        s.paragraph_format.space_after = Pt(6)
    for name, size in [('Heading 1', 17), ('Heading 2', 13), ('Heading 3', 11.5)]:
        s = doc.styles[name]; s.font.size = Pt(size)
        s.font.color.rgb = RGBColor.from_string('19354B')
        s.paragraph_format.space_before = Pt(15); s.paragraph_format.space_after = Pt(8)
        s.paragraph_format.keep_with_next = True
    doc.styles['Title'].font.size = Pt(26)
    doc.styles['Subtitle'].font.size = Pt(15)
    for name in ['Source Code', 'Verbatim Char']:
        if name in doc.styles:
            doc.styles[name].font.name = 'DejaVu Sans Mono'
            doc.styles[name].font.size = Pt(8)
    p = section.header.paragraphs[0]
    r = p.add_run('TAMU  |  影响函数 · 遗忘 · 可微调度 · 样本重加权')
    r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string('526675')
    p = section.footer.paragraphs[0]
    p.alignment = 1
    r = p.add_run('算法学习讲义  |  '); r.font.size = Pt(8)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); p._p.append(field)
    file = QA/'reference.docx'; doc.save(file)
    return file


def build_docx():
    command=['pandoc', str(OUT/'算法精读与推导讲义.md'), '--from=markdown+tex_math_dollars',
             '--to=docx', '--standalone',
             '--resource-path='+str(OUT), '--reference-doc='+str(reference()),
             '-o', str(OUT/(STEM+'.docx'))]
    result=subprocess.run(command, capture_output=True, text=True, check=True)
    if result.stderr.strip():
        print(result.stderr)
    assert 'Could not convert' not in result.stderr
    path=OUT/(STEM+'.docx')
    doc=Document(path)
    headings=[p for p in doc.paragraphs if p.style.name=='Heading 1']
    headings[0].paragraph_format.page_break_before=True
    title=headings[0].insert_paragraph_before('导读目录', style='Heading 2')
    for index,heading in enumerate(headings, 1):
        anchor=f'guide_ch_{index}'
        begin=OxmlElement('w:bookmarkStart');begin.set(qn('w:id'), str(900+index));begin.set(qn('w:name'), anchor)
        end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'), str(900+index))
        heading._p.insert(1, begin);heading._p.append(end)
        p=headings[0].insert_paragraph_before()
        p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.05
        link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'), anchor)
        run=OxmlElement('w:r');pr=OxmlElement('w:rPr')
        size=OxmlElement('w:sz');size.set(qn('w:val'), '19');pr.append(size)
        color=OxmlElement('w:color');color.set(qn('w:val'), '2460A6');pr.append(color)
        run.append(pr);text=OxmlElement('w:t');text.text=f'{index:02d}  {heading.text}'
        run.append(text);link.append(run);p._p.append(link)
    proportions=[[.10,.26,.37,.27], [.10,.52,.38], [.19,.40,.41],
                 [.12,.22,.22,.44], [.32,.68], [.28,.22,.25,.25], [.14,.47,.39]]
    # Rebuild Pandoc tables with python-docx's standard table skeleton; retain
    # the original paragraph XML so inline Office Math and links stay editable.
    original_tables=list(doc.tables)
    for old in original_tables:
        table=doc.add_table(rows=len(old.rows), cols=len(old.columns))
        for oldrow,newrow in zip(old.rows,table.rows):
            for oldcell,newcell in zip(oldrow.cells,newrow.cells):
                for p in list(newcell._tc.findall(qn('w:p'))):newcell._tc.remove(p)
                for p in oldcell._tc.findall(qn('w:p')):newcell._tc.append(deepcopy(p))
                for p in newcell.paragraphs:p.style='Normal'
        old._tbl.addprevious(table._tbl)
        old._tbl.getparent().remove(old._tbl)
    for table,weights in zip(doc.tables,proportions):
        table.style='Table Grid'
        table.autofit=False
        total=6.74
        tblw=table._tbl.tblPr.find(qn('w:tblW'))
        tblw.set(qn('w:type'), 'dxa');tblw.set(qn('w:w'),str(round(total*1440)))
        for col,fraction in zip(table.columns, weights):col.width=Inches(total*fraction)
        header=table.rows[0]
        repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'), 'true')
        header._tr.get_or_add_trPr().append(repeat)
        for index,row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell,fraction in zip(row.cells,weights):
                cell.width=Inches(total*fraction)
                if index == 0:
                    fill=OxmlElement('w:shd');fill.set(qn('w:fill'), 'EAF3F4')
                    cell._tc.get_or_add_tcPr().append(fill)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(3)
                    p.paragraph_format.line_spacing=1.1
                    for r in p.runs:
                        r.font.size=Pt(9)
    doc.core_properties.title='TAMU 算法精读与推导讲义'
    doc.core_properties.subject='先修文献链接、数学推导、算法边界及项目实现对应'
    doc.core_properties.author='task_aware_machine_unlearning 项目学习材料'
    doc.save(path)


def render():
    path=OUT/(STEM+'.docx')
    command=['soffice', '-env:UserInstallation=file:///tmp/tamu-algorithm-guide-lo',
             '--headless', '--convert-to', 'pdf', '--outdir', str(OUT), str(path)]
    result=subprocess.run(command, capture_output=True, text=True, timeout=110,
                          env={**os.environ, 'XDG_CACHE_HOME':'/workspace/.cache'})
    print(result.stdout.strip())
    if result.returncode or not path.with_suffix('.pdf').exists():
        raise RuntimeError(result.stdout+'\n'+result.stderr)


def check_evidence():
    # All four data-derived rows in the narrative must match committed records.
    import csv
    with (ROOT/'research/cnn_temperature_response/results.csv').open(encoding='utf-8-sig') as f:
        rows=list(csv.DictReader(f))
    values={'original': (588.473549, -.827542, -6.778363),
            'unlearned20':(680.181358, -.844525, -7.007316),
            'retrained':(680.180885, -.844474, -7.006892),
            'tamu20':(564.325866, -.290330, -2.371888)}
    for model, expected in values.items():
        zero=next(r for r in rows if r['model']==model and float(r['sigma_level'])==0)
        plus=next(r for r in rows if r['model']==model and float(r['sigma_level'])==1)
        actual=[float(zero['cost_total']), float(plus['forecast_total_delta_mw']), float(plus['cost_delta'])]
        assert np.allclose(actual, expected, rtol=0, atol=5.1e-7), (model, actual, expected)
    with (ROOT/'research/cnn_cost_attribution/paired_model_costs.csv').open(encoding='utf-8-sig') as f:
        paired=list(csv.DictReader(f))
    nz=[r for r in paired if float(r['sigma_level'])!=0]
    assert len(nz)==108 and sum(float(r['tamu_minus_direct']) < -1e-4 for r in nz)==106
    counter=next(r for r in nz if int(r['test_index'])==1686 and float(r['sigma_level'])==-1)
    assert abs(float(counter['tamu_minus_direct'])-60.929210)<5.1e-7
    manifest=json.loads((OUT/'source_manifest.json').read_text())
    for item in manifest['sources']:
        if 'file' in item:
            assert sha(ROOT/item['file'])==item['sha256'], item
    return dict(reference_rows_verified=4, nonzero_cohort_scenarios_verified=108,
                tamu_below_direct=106, counterexample_verified=True,
                source_hashes_verified=True)


def inspect():
    from lxml import etree
    docx=OUT/(STEM+'.docx')
    ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math',
        'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    with ZipFile(docx) as z:
        assert z.testzip() is None
        xml=etree.fromstring(z.read('word/document.xml'))
        math_count=len(xml.findall('.//m:oMath', ns))
        assert math_count>100, math_count
        relations=etree.fromstring(z.read('word/_rels/document.xml.rels'))
        urls=[r.get('Target') for r in relations if r.get('Type','').endswith('/hyperlink')]
        for expected in ['https://arxiv.org/abs/1703.04730',
                         'https://arxiv.org/abs/1703.00443',
                         'https://arxiv.org/abs/1910.12430',
                         'https://arxiv.org/abs/1703.04529',
                         'https://arxiv.org/abs/1911.03030']:
            assert expected in urls
        assert len(urls)>30
        raw_text=''.join(xml.itertext())
        assert '\\begin{' not in raw_text and '\\frac{' not in raw_text
    pages=[];previews=[];outside=[]
    with fitz.open(OUT/(STEM+'.pdf')) as pdf:
        alltext='\n'.join(p.get_text() for p in pdf)
        assert '没有取得目标论文' in alltext.replace('\n','')
        assert alltext.count('\ufffd')==0
        for i,page in enumerate(pdf):
            text=page.get_text()
            pages.append(dict(page=i+1, characters=len(text.strip())))
            assert len(text.strip())>40, (i+1, text)
            for block in page.get_text('dict')['blocks']:
                if block['type']!=0:continue
                for line in block['lines']:
                    for span in line['spans']:
                        box=fitz.Rect(span['bbox'])
                        if box.x0<-.8 or box.y0<-.8 or box.x1>page.rect.width+.8 or box.y1>page.rect.height+.8:
                            outside.append(dict(page=i+1,text=span['text'],bbox=list(box)))
            p=QA/f'page-{i+1:02d}.png'
            page.get_pixmap(matrix=fitz.Matrix(1,1), alpha=False).save(p)
            previews.append(p)
        (QA/'rendered-text.txt').write_text(alltext)
    assert not outside, outside
    width=285;height=410
    for start in range(0, len(previews), 12):
        chunk=previews[start:start+12];rows=(len(chunk)+2)//3
        contact=Image.new('RGB', (width*3, (height+25)*rows), '#DEE4E8')
        draw=ImageDraw.Draw(contact)
        for index,path in enumerate(chunk):
            with Image.open(path) as im:
                thumb=ImageOps.contain(im, (width-10, height-5))
                x=(index%3)*width+(width-thumb.width)//2
                y=(index//3)*(height+25)
                contact.paste(thumb, (x,y));draw.text(((index%3)*width+7,y+height),f'Page {start+index+1}', fill='black')
        contact.save(QA/f'contact-{start//12+1:02d}.png')
    result=dict(docx_native_equations=math_count, distinct_hyperlink_targets=len(set(urls)),
                pdf_pages=len(pages), page_details=pages, out_of_page_text=outside,
                primary_paper_full_text_obtained=False,
                source_limitation_disclosed=True, evidence=check_evidence(),
                teaching_checks=json.loads((OUT/'derivation_checks.json').read_text()))
    (OUT/'document_validation.json').write_text(json.dumps(result, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps({k: v for k,v in result.items() if k not in ['page_details','teaching_checks']}, ensure_ascii=False, indent=2))


def export():
    sums=[]
    for path in sorted(OUT.rglob('*')):
        if path.is_file() and path.name!='SHA256SUMS':
            sums.append(f'{sha(path)}  {path.relative_to(OUT)}')
    (OUT/'SHA256SUMS').write_text('\n'.join(sums)+'\n')
    archive=ROOT/'exports/tamu_algorithm_guide.zip'
    with ZipFile(archive, 'w', ZIP_DEFLATED) as z:
        for path in sorted(OUT.rglob('*')):
            if path.is_file():z.write(path, Path('TAMU算法精读与推导讲义')/path.relative_to(OUT))
    with ZipFile(archive) as z:assert z.testzip() is None
    archive.with_suffix('.zip.sha256').write_text(f'{sha(archive)}  {archive.name}\n')
    print('ZIP:', archive.name, archive.stat().st_size, sha(archive))


def main():
    QA.mkdir(parents=True, exist_ok=True)
    subprocess.run([sys.executable, str(OUT/'verify_derivations.py')], check=True,
                   stdout=(QA/'teaching-checks.log').open('w'))
    assets();build_docx();render();inspect();export()


if __name__=='__main__':
    main()
