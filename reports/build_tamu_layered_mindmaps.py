"""Build layered learning maps without claiming access to the paper full text."""
from copy import deepcopy
from datetime import datetime
import hashlib
import html
import importlib.util
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import unicodedata
import xml.etree.ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

import fitz
from PIL import Image, ImageDraw, ImageOps
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'reports/tamu_layered_mindmaps'
FIG = OUT/'figures'
QA = Path('/workspace/.cloud-setup/task-aware/tamu-mindmap-qa')
BASE = '960e7505f307cc4ac1cae17de655176cf29d7bd7'
FONT = 'Noto Sans CJK SC'
COLORS = ['#2460A6','#167D8D','#8C5A19','#5F52A3','#1D7562','#B2533B','#7856A0','#526675']
DISCLAIMER = '学习分层（非原文章节号）｜原文全文未取得；依据作者摘要、代码与标准推导整理'

spec=importlib.util.spec_from_file_location('map_content', OUT/'content.py')
content=importlib.util.module_from_spec(spec);spec.loader.exec_module(content)
LAYERS=content.LAYERS;REFS=content.REFERENCES


def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def q(text):return json.dumps(text, ensure_ascii=False)


def wrap(text, limit=32):
    result=[]
    for paragraph in text.split('\n'):
        line='';width=0
        for char in paragraph:
            w=2 if unicodedata.east_asian_width(char) in 'WF' else 1
            if width+w>limit and line:
                result.append(line);line='';width=0
            line+=char;width+=w
        result.append(line)
    return '\n'.join(result)


def graph(name, title, branches, color='#19354B'):
    source=['digraph study {',
        'graph [rankdir=LR, bgcolor="white", nodesep=0.18, ranksep=0.38, pad=0.22, splines=polyline, outputorder=edgesfirst];',
        f'graph [label={q(DISCLAIMER)}, labelloc=b, fontsize=11, fontname={q(FONT)}, fontcolor="#526675"];',
        f'node [shape=box, style="rounded,filled", fontname={q(FONT)}, fontsize=18, margin="0.17,0.12", color="#B5C6D0", fillcolor="white", penwidth=1.2];',
        'edge [arrowhead=none, color="#8AA3B1", penwidth=1.4];',
        f'root [label={q(wrap(title,28))}, fontsize=22, fillcolor={q(color)}, fontcolor="white", color={q(color)}];']
    ids=['root']
    for i,branch in enumerate(branches):
        branch_id=f'b{i}';ids.append(branch_id)
        branch_color=branch.get('color',color)
        source.append(f'{branch_id} [label={q(wrap(branch["label"],28))}, color={q(branch_color)}, fillcolor="#EFF5F8", penwidth=1.8];')
        source.append(f'root -> {branch_id} [color={q(branch_color)}];')
        for j,child in enumerate(branch['children']):
            node=f'b{i}_l{j}';ids.append(node)
            source.append(f'{node} [label={q(wrap(child["label"],36))}, tooltip={q(child.get("detail", ""))}];')
            source.append(f'{branch_id} -> {node} [color={q(branch_color)}];')
    source.append('}')
    dot=FIG/(name+'.dot');dot.write_text('\n'.join(source)+'\n')
    for extension in ['svg','png','pdf']:
        cmd=['dot', '-T'+extension, str(dot), '-o', str(FIG/(name+'.'+extension))]
        if extension=='png':cmd.insert(1,'-Gdpi=160')
        subprocess.run(cmd,check=True,capture_output=True)
    root=ET.parse(FIG/(name+'.svg')).getroot()
    node_count=sum(1 for el in root.iter() if el.get('class')=='node')
    assert node_count==len(ids), (name,node_count,len(ids))
    return dict(name=name,nodes=node_count,title=title)


def make_maps():
    overview=[]
    for layer,color in zip(LAYERS,COLORS):
        overview.append(dict(label='第'+layer['id']+'层\n'+layer['title'],color=color,
                             children=[dict(label=layer['summary'],detail=layer['key'])]))
    result=[graph('00_overview','Task-Aware\nMachine Unlearning\n任务感知机器遗忘',overview)]
    for layer,color in zip(LAYERS,COLORS):
        result.append(graph(layer['id']+'_layer', '第'+layer['id']+'层\n'+layer['title'],layer['groups'],color))
    return result


def make_markdown():
    lines=['# Task-Aware Machine Unlearning 分层阅读与导图','',
        '**'+DISCLAIMER+'。本套图是学习结构，不能冒充已逐章核对的论文目录。**','',
        '目标论文：[arXiv 2308.14412](https://arxiv.org/abs/2308.14412)。', '',
        '![总导图](figures/00_overview.png)','',
        '阅读顺序：为什么删 → 删什么 → 怎样直接删 → 怎样保预测性能 → 怎样定义任务成本 → 怎样按成本调权 → 怎样验证 → 能下什么结论。','']
    for layer in LAYERS:
        lines.extend(['## 第'+layer['id']+'层：'+layer['title'],'',layer['subtitle'],'',
                      '**来源范围：**'+layer['source'],'',layer['explanation'],'',
                      '![分层详图](figures/'+layer['id']+'_layer.png)',''])
        if 'formula' in layer:lines.extend(['**关键关系（本讲义记号）：** '+layer['formula'],''])
        for g in layer['groups']:
            lines.extend(['### '+g['label'],''])
            for child in g['children']:
                lines.extend(['- **'+child['label'].replace('\n','：')+'**：'+child['detail']])
            lines.append('')
        lines.extend(['**读懂标志：**'+layer['key'],'', '**代码位置：**'])
        for path in layer['code']:
            file=path.split(':')[0]
            lines.append('- ['+path+'](https://github.com/syh-Mason/task_aware_machine_unlearning/blob/'+BASE+'/'+file+')')
        lines.extend(['','**先修来源：**'+'；'.join('['+REFS[r][0]+']('+REFS[r][1]+')' for r in layer['refs']), ''])
    lines.extend(['## 原论文与扩展实验的边界','',
        '第07层中的温度单点、多日期、节点隔离/复位/Shapley 是本项目后续扩展，不是原论文已经报告的实验。原作者摘要确认的模型与研究范围、当前代码实现、标准数学推导和扩展实测结果分别标记。', '',
        '原论文的准确章节标题、编号、定理假设、证明、图表与发表版差异，需要取得 PDF 后逐项核验。本图没有给未知内容填入虚构编号或数值。',''])
    (OUT/'分层阅读说明.md').write_text('\n'.join(lines))
    # Mermaid offers an additional text-editable overview, without pretending
    # that its learning layers are the source paper's numbered sections.
    m=['mindmap','  root((Task-Aware Machine Unlearning))']
    for layer in LAYERS:
        m.append('    '+layer['id']+' '+layer['title'])
        m.extend('      '+g['label'] for g in layer['groups'])
    (OUT/'总导图.mmd').write_text('\n'.join(m)+'\n')


def make_freemind():
    root=ET.Element('map', {'version':'1.0.1'})
    node=ET.SubElement(root,'node',{'TEXT':'Task-Aware Machine Unlearning｜学习分层（非原文章节）'})
    ET.SubElement(node,'node',{'TEXT':DISCLAIMER})
    for layer,color in zip(LAYERS,COLORS):
        child=ET.SubElement(node,'node',{'TEXT':layer['id']+' '+layer['title'],'COLOR':color})
        for g in layer['groups']:
            sub=ET.SubElement(child,'node',{'TEXT':g['label']})
            for l in g['children']:
                leaf=ET.SubElement(sub,'node',{'TEXT':l['label'].replace('\n','：')})
                ET.SubElement(leaf,'node',{'TEXT':l['detail'],'FOLDED':'true'})
    ET.indent(root)
    ET.ElementTree(root).write(OUT/'TAMU_可编辑导图.mm',encoding='utf-8',xml_declaration=True)


def make_html(maps):
    options=[];panes=[];notes=[]
    for i,m in enumerate(maps):
        label='总导图' if i==0 else LAYERS[i-1]['id']+' '+LAYERS[i-1]['title']
        options.append('<button class="tab'+(' selected' if i==0 else '')+'" data-index="'+str(i)+'">'+html.escape(label)+'</button>')
        svg=(FIG/(m['name']+'.svg')).read_text()
        svg=svg[svg.index('<svg'):]
        panes.append('<section class="pane'+(' visible' if i==0 else '')+'" data-index="'+str(i)+'">'+svg+'</section>')
        if i==0:
            text='<h2>先看主线</h2><p>为什么删 → 删什么 → 怎样删 → 怎样保性能 → 怎样算成本 → 怎样按成本调权 → 怎样验证 → 怎样解释。</p>'
            text+='<p>选择左侧层级查看详图。滚轮缩放，拖动平移；点击“复位”恢复全图。下方保留各层的完整解释。</p>'
        else:
            layer=LAYERS[i-1]
            text='<h2>'+html.escape(layer['title'])+'</h2><p>'+html.escape(layer['explanation'])+'</p>'
            if 'formula' in layer:text+='<p class="formula">'+html.escape(layer['formula'])+'</p>'
            for group in layer['groups']:
                text+='<h3>'+html.escape(group['label'])+'</h3><ul>'
                for leaf in group['children']:
                    text+='<li><b>'+html.escape(leaf['label'].replace('\n','：'))+'</b>：'+html.escape(leaf['detail'])+'</li>'
                text+='</ul>'
            text+='<p><b>读懂标志：</b>'+html.escape(layer['key'])+'</p><p><b>来源：</b>'+html.escape(layer['source'])+'</p>'
            text+='<p>'+' · '.join('<a href="'+REFS[r][1]+'" target="_blank" rel="noopener">'+html.escape(REFS[r][0])+'</a>' for r in layer['refs'])+'</p>'
        notes.append('<article class="notes'+(' visible' if i==0 else '')+'" data-index="'+str(i)+'">'+text+'</article>')
    document='''<!DOCTYPE html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>TAMU 分层交互导图</title><style>
*{box-sizing:border-box}body{margin:0;background:#f3f6f8;color:#19354b;font:15px/1.65 system-ui,"Noto Sans CJK SC",sans-serif}
header{padding:16px 22px;background:#19354b;color:white}header h1{margin:0;font-size:23px}.limit{margin:6px 0 0;color:#f1d4a0;font-size:13px}
.layout{display:grid;grid-template-columns:230px minmax(0,1fr);min-height:75vh}nav{padding:16px;background:#e8f0f4}.tab{display:block;width:100%;text-align:left;border:0;border-radius:7px;padding:12px;margin-bottom:6px;background:transparent;color:#19354b;font-size:15px;cursor:pointer}.tab:hover,.tab.selected{background:#19354b;color:white}
main{min-width:0;padding:16px}.toolbar{display:flex;gap:8px;flex-wrap:wrap;align-items:center;margin-bottom:10px}.toolbar button{padding:8px 14px;border:1px solid #a9bdca;border-radius:6px;background:white;cursor:pointer}input{padding:8px;border:1px solid #a9bdca;border-radius:6px;min-width:210px}.status{font-size:13px;color:#526675}
.canvas{height:71vh;min-height:480px;border:1px solid #cad7df;background:white;border-radius:9px;overflow:hidden}.pane{display:none;width:100%;height:100%;touch-action:none}.pane.visible{display:block}.pane svg{width:100%;height:100%;display:block;cursor:grab}.pane svg:active{cursor:grabbing}.node.matched polygon,.node.matched path{stroke:#e08022!important;stroke-width:3px!important;fill:#fff0c4!important}
.notes{display:none;background:white;padding:22px;margin-top:16px;border:1px solid #d3dfe6;border-radius:9px}.notes.visible{display:block}.notes h2{margin-top:0}.notes li{margin-bottom:7px}.formula{padding:10px;background:#edf5f6;border-left:3px solid #167d8d}a{color:#2460a6}
@media(max-width:760px){.layout{display:block}nav{display:flex;overflow:auto;padding:8px}.tab{min-width:155px;width:auto;margin:0 5px 0 0;padding:8px}.canvas{height:63vh}.limit{font-size:12px}}@media print{nav,.toolbar{display:none}.layout{display:block}.canvas{height:80vh}.notes{page-break-before:always}}
</style><header><h1>Task-Aware Machine Unlearning｜分层交互导图</h1><p class="limit">DISCLAIMER_PLACEHOLDER</p></header>
<div class="layout"><nav>OPTIONS_PLACEHOLDER</nav><main><div class="toolbar"><button id="plus">放大＋</button><button id="minus">缩小－</button><button id="reset">复位</button><input id="search" placeholder="搜索当前图，例如 Cost / Hessian" aria-label="搜索当前图"><span class="status">滚轮缩放 · 拖动平移 · 图下有逐项解释</span></div><div class="canvas">PANES_PLACEHOLDER</div>NOTES_PLACEHOLDER</main></div>
<script>
let current=0;const initial=new Map();let dragging=null;
function svg(){return document.querySelector('.pane.visible svg')}
document.querySelectorAll('.pane svg').forEach(s=>{initial.set(s,s.getAttribute('viewBox'));s.addEventListener('wheel',e=>{e.preventDefault();zoom(e.deltaY<0?.88:1.14,e.clientX,e.clientY)},{passive:false});s.addEventListener('pointerdown',e=>{if(e.button!==0)return;dragging={s,x:e.clientX,y:e.clientY,v:s.viewBox.baseVal.x,w:s.viewBox.baseVal.y};s.setPointerCapture(e.pointerId)});s.addEventListener('pointermove',e=>{if(!dragging||dragging.s!==s)return;const r=s.getBoundingClientRect(),v=s.viewBox.baseVal;s.setAttribute('viewBox',`${dragging.v-(e.clientX-dragging.x)*v.width/r.width} ${dragging.w-(e.clientY-dragging.y)*v.height/r.height} ${v.width} ${v.height}`)});s.addEventListener('pointerup',()=>dragging=null);s.addEventListener('pointercancel',()=>dragging=null)});
function zoom(f,cx,cy){let s=svg(),v=s.viewBox.baseVal,r=s.getBoundingClientRect();let fx=cx===undefined?.5:Math.max(0,Math.min(1,(cx-r.left)/r.width)),fy=cy===undefined?.5:Math.max(0,Math.min(1,(cy-r.top)/r.height));const w=v.width*f,h=v.height*f;if(w<40||w>15000)return;s.setAttribute('viewBox',`${v.x+(v.width-w)*fx} ${v.y+(v.height-h)*fy} ${w} ${h}`)}
function search(){const text=document.getElementById('search').value.trim().toLowerCase();svg().querySelectorAll('.node').forEach(n=>n.classList.toggle('matched',!!text&&n.textContent.toLowerCase().includes(text)))}
document.querySelectorAll('.tab').forEach(b=>b.addEventListener('click',()=>{current=Number(b.dataset.index);document.querySelectorAll('.tab').forEach(x=>x.classList.toggle('selected',x===b));document.querySelectorAll('.pane,.notes').forEach(x=>x.classList.toggle('visible',Number(x.dataset.index)===current));search()}));
document.getElementById('plus').onclick=()=>zoom(.8);document.getElementById('minus').onclick=()=>zoom(1.25);document.getElementById('reset').onclick=()=>svg().setAttribute('viewBox',initial.get(svg()));document.getElementById('search').addEventListener('input',search);
</script></html>'''
    document=document.replace('DISCLAIMER_PLACEHOLDER',html.escape(DISCLAIMER)).replace('OPTIONS_PLACEHOLDER',''.join(options)).replace('PANES_PLACEHOLDER',''.join(panes)).replace('NOTES_PLACEHOLDER',''.join(notes))
    (OUT/'TAMU_交互导图.html').write_text(document)
    script=document.split('<script>')[1].split('</script>')[0]
    (QA/'interactive.js').write_text(script)


def map_pdf(maps):
    pdf=fitz.open()
    for item in maps:
        page=pdf.new_page(width=1190.55,height=841.89)
        with fitz.open(FIG/(item['name']+'.pdf')) as source:
            page.show_pdf_page(fitz.Rect(30,30,1160.55,811.89),source,0,keep_proportion=True)
    path=OUT/'TAMU_分层导图合集.pdf';pdf.save(path,deflate=True);pdf.close()


def font(run,size=None,bold=None):
    run.font.name=FONT;run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
    if size:run.font.size=Pt(size)
    if bold is not None:run.bold=bold


def make_word():
    doc=Document();s=doc.sections[0]
    s.page_width=Inches(8.2677);s.page_height=Inches(11.6929)
    s.left_margin=s.right_margin=Inches(.68);s.top_margin=s.bottom_margin=Inches(.65)
    for name,size in [('Normal',10.5),('Title',24),('Heading 1',17),('Heading 2',12.5),('Caption',9)]:
        st=doc.styles[name];st.font.name=FONT;st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT);st.font.size=Pt(size)
    doc.styles['Normal'].paragraph_format.line_spacing=1.2
    doc.styles['Normal'].paragraph_format.space_after=Pt(7)
    font(s.header.paragraphs[0].add_run('TAMU｜按研究逻辑分层阅读（代码核对版）'),8)
    p=s.footer.paragraphs[0];p.alignment=1;font(p.add_run('学习导图  |  '),8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
    doc.add_heading('Task-Aware Machine Unlearning\n分层阅读与导图',0)
    doc.add_paragraph('2026-10-21｜一张总图 · 八张详图 · 逐层解释')
    doc.add_paragraph(DISCLAIMER+'。',style='Subtitle')
    doc.add_paragraph('这套图按学习顺序拆分作者方法，并不是原论文目录的转录。原文站点仍返回代理403，准确章节号、公式号、定理与证明待取得全文后核验。所有温度扰动及多日期结果单列为你的扩展实验。')
    for layer in LAYERS:
        doc.add_paragraph(layer['id']+'  '+layer['title']+'：'+layer['summary'])
    doc.add_paragraph('用法：先看总图，再按01→08阅读。每层先回答“解决什么问题”，再顺着三条分支读。图中的详细节点解释同时保存在 Markdown 和可离线交互 HTML 中。')
    doc.add_page_break();doc.add_heading('总导图：八层怎样连起来',1)
    with Image.open(FIG/'00_overview.png') as image:
        w,h=image.size;maxw,maxh=6.85,8.15;scale=min(maxw/w,maxh/h)
        doc.add_picture(str(FIG/'00_overview.png'),width=Inches(w*scale))
    doc.add_paragraph('学习层次不是原文章节；SVG 与A3矢量PDF可以放大查看。',style='Caption')
    for layer in LAYERS:
        doc.add_page_break();doc.add_heading('第'+layer['id']+'层｜'+layer['title'],1)
        doc.add_paragraph(layer['subtitle'])
        with Image.open(FIG/(layer['id']+'_layer.png')) as image:
            w,h=image.size;scale=min(6.85/w,4.65/h)
            doc.add_picture(str(FIG/(layer['id']+'_layer.png')),width=Inches(w*scale))
        doc.add_paragraph(layer['source'],style='Caption')
        doc.add_paragraph(layer['explanation'])
        if 'formula' in layer:
            p=doc.add_paragraph();font(p.add_run(layer['formula']),10)
        p=doc.add_paragraph();font(p.add_run('读懂标志：'),bold=True);p.add_run(layer['key'])
        doc.add_paragraph('代码：'+'；'.join(layer['code']))
        doc.add_paragraph('先修材料：'+'；'.join(REFS[r][0] for r in layer['refs']))
    doc.add_page_break();doc.add_heading('来源与使用边界',1)
    doc.add_paragraph('本套学习分层不是对原论文全部章节的已核验分析。作者摘要、代码实现、标准数学推导、你的扩展实验具有不同来源，图中已注明。以下论文全文链接不代表本次成功取得正文。')
    for key,(title,url) in REFS.items():
        doc.add_paragraph(key+'｜'+title+'\n'+url)
    doc.add_paragraph('原文 PDF 到位后，应核对章节、问题定义、权重变量约定、评分锚点、梯度存在定理的完整假设和证明，以及发表版实验图表。当前没有用未知原文数据填充这些内容。')
    doc.core_properties.title='TAMU 分层阅读与导图（代码核对版）'
    doc.core_properties.author='task_aware_machine_unlearning 项目学习材料'
    path=OUT/'TAMU_分层阅读说明.docx';doc.save(path)
    command=['soffice','-env:UserInstallation=file:///tmp/tamu-layered-maps-lo','--headless','--convert-to','pdf','--outdir',str(OUT),str(path)]
    r=subprocess.run(command,capture_output=True,text=True,timeout=100,env={**os.environ,'XDG_CACHE_HOME':'/workspace/.cache'})
    if r.returncode or not path.with_suffix('.pdf').exists():raise RuntimeError(r.stdout+'\n'+r.stderr)
    print(r.stdout.strip())


def verify(maps):
    assert len(LAYERS)==8 and len(maps)==9
    assert all(len(l['groups'])==3 and all(len(g['children'])==3 for g in l['groups']) for l in LAYERS)
    bundle=fitz.open(OUT/'TAMU_分层导图合集.pdf');assert len(bundle)==9
    counts=[]
    for i,p in enumerate(bundle):
        text=p.get_text();assert len(text)>100 and '\ufffd' not in text
        counts.append(len(text))
    bundle.close()
    previews=[];bounds=[]
    with fitz.open(OUT/'TAMU_分层阅读说明.pdf') as pdf:
        for i,p in enumerate(pdf):
            text=p.get_text();assert len(text.strip())>50 and '\ufffd' not in text
            path=QA/f'word-page-{i+1:02d}.png';p.get_pixmap(matrix=fitz.Matrix(.75,.75),alpha=False).save(path);previews.append(path)
            for block in p.get_text('dict')['blocks']:
                if block['type']!=0:continue
                for line in block['lines']:
                    for span in line['spans']:
                        b=fitz.Rect(span['bbox'])
                        if b.x0<-.8 or b.y0<-.8 or b.x1>p.rect.width+.8 or b.y1>p.rect.height+.8:bounds.append(dict(page=i+1,text=span['text']))
        wordpages=len(pdf)
    assert not bounds,bounds
    cols=3;tw=300;th=425
    sheet=Image.new('RGB',(cols*tw,((len(previews)+cols-1)//cols)*(th+24)),'#DDE4EA');draw=ImageDraw.Draw(sheet)
    for i,p in enumerate(previews):
        with Image.open(p) as im:
            thumb=ImageOps.contain(im,(tw-10,th-4));x=(i%cols)*tw+(tw-thumb.width)//2;y=(i//cols)*(th+24)
            sheet.paste(thumb,(x,y));draw.text(((i%cols)*tw+8,y+th),'Page '+str(i+1),fill='black')
    sheet.save(QA/'word-contact.png')
    data=json.loads((OUT/'mindmap.json').read_text());assert data['primary_paper_full_text_obtained'] is False
    fm=ET.parse(OUT/'TAMU_可编辑导图.mm');assert len(fm.findall('./node/node'))==9
    with ZipFile(OUT/'TAMU_分层阅读说明.docx') as z:assert z.testzip() is None
    node=shutil.which('node')
    if node:subprocess.run([node,'--check',str(QA/'interactive.js')],check=True,capture_output=True)
    # Descriptive numbers come only from saved extension results.
    import csv
    with (ROOT/'research/cnn_temperature_response/results.csv').open(encoding='utf-8-sig') as f:rows=list(csv.DictReader(f))
    plus=next(r for r in rows if r['model']=='original' and float(r['sigma_level'])==1)
    assert abs(float(plus['forecast_total_delta_mw'])+.827542)<5.1e-7
    assert abs(float(plus['cost_delta'])+6.778363)<5.1e-7
    with (ROOT/'research/cnn_cost_attribution/paired_model_costs.csv').open(encoding='utf-8-sig') as f:rows=list(csv.DictReader(f))
    nz=[r for r in rows if float(r['sigma_level'])!=0]
    assert len(nz)==108 and sum(float(r['tamu_minus_direct']) < -1e-4 for r in nz)==106
    counter=next(r for r in nz if int(r['test_index'])==1686 and float(r['sigma_level'])==-1)
    assert abs(float(counter['tamu_minus_direct'])-60.929210)<5.1e-7
    result=dict(learning_layers=8,maps=maps,map_pdf_pages=9,map_page_characters=counts,
                word_pdf_pages=wordpages,out_of_page_text=bounds,
                primary_paper_full_text_obtained=False,extension_results_verified=True,
                all_figures_have_source_disclaimer=True,interactive_html_self_contained=True,
                interactive_javascript_syntax_checked=bool(node),
                html_runtime_browser_tested=False,
                interactive_browser_limitation='No browser runtime test claimed; standalone files and script syntax checked.')
    (OUT/'validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(result,ensure_ascii=False,indent=2))


def export():
    files=[p for p in sorted(OUT.rglob('*')) if p.is_file() and '__pycache__' not in p.parts and p.name!='SHA256SUMS']
    (OUT/'SHA256SUMS').write_text('\n'.join(sha(p)+'  '+str(p.relative_to(OUT)) for p in files)+'\n')
    archive=ROOT/'exports/tamu_layered_mindmaps.zip'
    with ZipFile(archive,'w',ZIP_DEFLATED) as z:
        for p in files+[OUT/'SHA256SUMS']:z.write(p,Path('TAMU分层导图')/p.relative_to(OUT))
    with ZipFile(archive) as z:assert z.testzip() is None
    archive.with_suffix('.zip.sha256').write_text(sha(archive)+'  '+archive.name+'\n')
    print('ZIP',archive.stat().st_size,sha(archive))


def main():
    FIG.mkdir(exist_ok=True);QA.mkdir(parents=True,exist_ok=True)
    data=dict(title='TAMU 分层阅读（非原文目录）',compiled_date='2026-10-21',
              disclaimer=DISCLAIMER,primary_paper_full_text_obtained=False,
              base_commit=BASE,references=REFS,layers=LAYERS)
    (OUT/'mindmap.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')
    retrieval=Path('/workspace/.cloud-setup/task-aware/mindmap-sources/retrieval.json')
    if retrieval.exists():records=json.loads(retrieval.read_text())
    else:records=json.loads((OUT/'source_manifest.json').read_text())['retrieval']
    sources={name:sha(ROOT/name) for name in ['eval_unlearn.py','eval_unchange.py','func_operation.py','utils/net.py','utils/optimization.py','conf/model/conv.yaml','research/cnn_temperature_response/results.csv','research/cnn_cost_attribution/paired_model_costs.csv']}
    (OUT/'source_manifest.json').write_text(json.dumps(dict(base_commit=BASE,retrieval=records,source_hashes=sources,primary_paper_full_text_obtained=False,builder_sha256=sha(Path(__file__)),content_sha256=sha(OUT/'content.py')),ensure_ascii=False,indent=2)+'\n')
    maps=make_maps();make_markdown();make_freemind();make_html(maps);map_pdf(maps);make_word();verify(maps);export()


if __name__=='__main__':main()
