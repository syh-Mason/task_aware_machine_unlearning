"""Build a 10–13 minute source-bounded Task-Aware algorithm PPT work draft.

No new explanatory diagrams are generated: the one author repository image is
used unchanged; equations are typeset, and optional figures get text briefs.
Full arXiv v2 paper and NSE guideline page are not available in this workspace.
"""
from copy import deepcopy
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import unicodedata
from zipfile import ZipFile, ZIP_DEFLATED

os.environ.setdefault('MPLCONFIGDIR', '/workspace/.cloud-setup/task-aware/matplotlib')
os.environ.setdefault('XDG_CACHE_HOME', '/workspace/.cache')

import fitz
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from PIL import Image, ImageDraw, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches as DInches, Pt as DPt

from task_aware_talk_content import SLIDES, SUPPLEMENT, PAPER, AUTHOR, MIT_SPEAK, MIT_SLIDES

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'reports/task_aware_paper_talk'
SOURCE = Path('/workspace/.cloud-setup/task-aware/task-aware-slides-sources')
QA = Path('/workspace/.cloud-setup/task-aware/task-aware-slides-qa')
STEM = 'TaskAware_论文算法解析_10-13分钟_待原文核验'
FONT = 'Noto Sans CJK SC'
NAVY, TEAL, ORANGE = '16354A', '087F8C', 'B75331'
TEXT, GRAY, LIGHT, PALE = '243D4E', '596B77', 'EAF3F6', 'FBF2EC'
BRANCH = 'slides/task-aware-paper-mit'
BASE = f'https://github.com/syh-Mason/task_aware_machine_unlearning/blob/{BRANCH}'


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def color(hexcolor):
    return RGBColor.from_string(hexcolor)


def rect(slide, x, y, w, h, fill, line=None, rounded=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if line:
        shape.line.color.rgb = color(line)
        shape.line.width = Pt(.7)
    else:
        shape.line.fill.background()
    if rounded:
        shape.adjustments[0] = .05
    return shape


def text(slide, value, x, y, w, h, size=22, bold=False, ink=TEXT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    for i, line in enumerate(value.split('\n')):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.space_after = Pt(5)
        p.space_before = Pt(0)
        p.line_spacing = 1.10
        r = p.add_run()
        r.text = line
        r.font.name, r.font.size = FONT, Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color(ink)
        pr = r._r.get_or_add_rPr()
        ea = OxmlElement('a:ea')
        ea.set('typeface', FONT)
        pr.append(ea)
    return box


def picture(slide, path, x, y, w, h):
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    width, height = iw * scale, ih * scale
    shape = slide.shapes.add_picture(str(path), Inches(x + (w - width) / 2),
                                     Inches(y + (h - height) / 2),
                                     width=Inches(width), height=Inches(height))
    return shape


def formula_assets():
    folder = OUT / 'assets/equations'
    folder.mkdir(parents=True, exist_ok=True)
    manifest = []
    for slide in SLIDES:
        for i, expression in enumerate(slide.get('equations', [])):
            expression = expression.replace(r'\bigg|_{0}', r'|_{0}')
            expression = re.sub(r'\\le(?![a-zA-Z])', r'\\leq', expression)
            expression = re.sub(r'\\ge(?![a-zA-Z])', r'\\geq', expression)
            stem = f'{slide["id"]}_{i+1:02d}'
            # Typeset directly to the text's extent, without editing source art.
            fig = plt.figure(figsize=(18, 1.8), facecolor='white')
            item = fig.text(0, .5, expression, fontsize=34, va='center', color='#' + NAVY)
            fig.canvas.draw()
            extent = item.get_window_extent(fig.canvas.get_renderer()).transformed(fig.dpi_scale_trans.inverted())
            extent.x0 -= .05
            extent.x1 += .05
            extent.y0 -= .05
            extent.y1 += .05
            for ext in ['png', 'svg']:
                fig.savefig(folder / (stem + '.' + ext), dpi=240, bbox_inches=extent, facecolor='white')
                if ext == 'svg':
                    path = folder / (stem + '.svg')
                    path.write_text('\n'.join(line.rstrip() for line in path.read_text().splitlines()) + '\n')
            plt.close(fig)
            manifest.append(dict(slide=slide['id'], equation_index=i + 1, latex=expression,
                                 provenance='Explanatory math restatement, not an extracted paper equation or figure',
                                 png='assets/equations/' + stem + '.png'))
    (OUT / 'equations.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n')


def header(slide, item, number):
    rect(slide, 0, 0, 13.333, .10, TEAL if item['seconds'] else GRAY)
    text(slide, 'TASK-AWARE  /  ' + item['section'], .55, .22, 8.8, .28, 11, True, TEAL)
    timing = f'{item["seconds"]} 秒' if item['seconds'] else '问答备用 · 不计主讲'
    text(slide, timing, 10.4, .22, 2.35, .28, 11, False, GRAY)
    text(slide, item['title'], .55, .72, 12.25, .65, 29, True, NAVY)
    rect(slide, .55, 6.37, 12.23, .57, LIGHT)
    text(slide, item['takeaway'], .76, 6.48, 11.80, .37, 18, True, TEAL)
    text(slide, item['source'], .55, 7.10, 11.82, .26, 8.5, False, GRAY)
    text(slide, f'{number:02d}', 12.47, 7.08, .4, .28, 10, False, GRAY)


def add_equations(slide, item, y=1.58, gap=.75, maxheight=.57):
    for i, _ in enumerate(item.get('equations', [])):
        path = OUT / 'assets/equations' / f'{item["id"]}_{i+1:02d}.png'
        shape = picture(slide, path, .85, y + i * gap, 11.60, maxheight)
        shape.name = f'{item["id"]}:equation:{i+1}'
        shape._element.nvPicPr.cNvPr.set('descr', item['equations'][i])


def add_lines(slide, lines, y, gap=.60, size=22, x=.85, w=11.55):
    for i, line in enumerate(lines):
        text(slide, line, x, y + i * gap, w, .53, size, False, TEXT)


def build_presentation():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    for number, item in enumerate(SLIDES, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = color('FFFFFF')
        if item['kind'] == 'cover':
            rect(slide, 0, 0, .16, 7.5, TEAL)
            text(slide, item['kicker'], .75, .65, 11.8, .4, 15, True, TEAL)
            text(slide, item['title'], .75, 1.5, 11.8, 1.4, 35, True, NAVY)
            text(slide, item['subtitle'], .78, 3.15, 11.70, .76, 19, False, GRAY)
            text(slide, item['lines'][0], .78, 4.13, 11.70, .6, 18, False, GRAY)
            text(slide, item['lines'][1], .78, 5.08, 11.70, .7, 22, True, TEAL)
            text(slide, '10–13 分钟算法汇报  |  12 张主讲 + 5 张备份  |  计划 11 分 50 秒',
                 .78, 6.20, 11.70, .4, 15, False, GRAY)
            text(slide, '工作稿：论文 v2 全文、原文图表及定理假设尚待核对。',
                 .78, 6.96, 11.70, .3, 11, False, ORANGE)
        else:
            header(slide, item, number)
            kind = item['kind']
            if kind == 'original':
                picture(slide, OUT / 'assets/original/author_repository_unlearning.png', .65, 1.55, 8.02, 4.55)
                rect(slide, 8.97, 1.72, 3.60, 4.18, LIGHT, rounded=True)
                for i, line in enumerate(item['lines']):
                    text(slide, line, 9.18, 2.0 + i * 1.18, 3.15, 1.0, 21, i == 0)
            elif kind in ['direct', 'pamu', 'weights', 'tamu', 'stages', 'kkt', 'scores', 'backup_if', 'backup_quadratic']:
                count = len(item.get('equations', []))
                add_equations(slide, item, gap=.75, maxheight=.58)
                bottom = 1.58 + count * .75 + .16
                add_lines(slide, item['lines'], bottom, gap=.55, size=21 if count == 3 else 22)
            elif kind == 'procedure':
                add_lines(slide, item['lines'], 1.65, gap=.62, size=22)
                add_equations(slide, item, y=4.42, gap=.77, maxheight=.53)
            elif kind == 'scope':
                for i, line in enumerate(item['lines']):
                    rect(slide, .80, 1.78 + i * 1.25, .08, .83, TEAL if i != 2 else ORANGE)
                    text(slide, line, 1.10, 1.9 + i * 1.25, 11.00, .9, 23)
            elif kind == 'closing':
                for i, line in enumerate(item['lines']):
                    text(slide, f'0{i+1}', .85, 1.84 + i * 1.25, .8, .70, 32, True, TEAL)
                    text(slide, line, 1.95, 1.9 + i * 1.25, 10.40, .9, 25, True, NAVY)
            elif kind == 'backup_constraints':
                add_equations(slide, item, y=1.5, gap=.64, maxheight=.49)
                add_lines(slide, item['lines'], 4.49, gap=.49, size=18)
            elif kind == 'backup_figures':
                add_lines(slide, item['lines'], 1.70, gap=.67, size=21)
                rect(slide, .90, 4.08, 5.55, 1.55, PALE, ORANGE, rounded=True)
                rect(slide, 6.85, 4.08, 5.55, 1.55, PALE, ORANGE, rounded=True)
                text(slide, '原文结果图 R1\n待 PDF · 图号 / 子图待核对', 1.15, 4.43, 5.05, .95, 20, True, ORANGE)
                text(slide, '原文对照图 R2\n待 PDF · 不用复现图顶替', 7.10, 4.43, 5.05, .95, 20, True, ORANGE)
            elif kind == 'backup_sources':
                add_lines(slide, item['lines'], 1.70, gap=.92, size=21)
                links = [('目标论文 v2', PAPER), ('作者实现', AUTHOR),
                         ('MIT How to Speak', MIT_SPEAK), ('MIT NSE Slideshow', MIT_SLIDES)]
                for i, (label, url) in enumerate(links):
                    shape = text(slide, label, .9 + (i % 2) * 6.05, 4.65 + (i // 2) * .56,
                                 5.7, .45, 18, True, TEAL)
                    shape.text_frame.paragraphs[0].runs[0].hyperlink.address = url
            else:
                raise ValueError(kind)
        note = f'{item["id"]}｜计划 {item["seconds"]} 秒\n\n{item["notes"]}\n\n来源：{item["source"]}'
        if item.get('equations'):
            note += '\n\n公式源文本：\n' + '\n'.join(item['equations'])
        slide.notes_slide.notes_text_frame.text = note
    prs.core_properties.title = 'Task-Aware Machine Unlearning｜10–13 分钟算法解析工作稿'
    prs.core_properties.subject = '作者实现核对；论文 v2 与实验原图尚待核验；MIT 风格适配'
    prs.core_properties.author = 'task_aware_machine_unlearning 项目研究材料'
    target = OUT / (STEM + '.pptx')
    prs.save(target)
    (OUT / 'slides.json').write_text(json.dumps(SLIDES, ensure_ascii=False, indent=2) + '\n')
    return target


def document_style(doc, title):
    s = doc.sections[0]
    s.page_width, s.page_height = DInches(8.2677), DInches(11.6929)
    s.left_margin = s.right_margin = DInches(.75)
    s.top_margin = s.bottom_margin = DInches(.70)
    for name, size in [('Normal', 10.5), ('Title', 24), ('Heading 1', 16), ('Heading 2', 12)]:
        st = doc.styles[name]
        st.font.name = FONT
        st.font.size = DPt(size)
        st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.2
    s.header.paragraphs[0].text = title
    footer = s.footer.paragraphs[0]
    footer.alignment = 1
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)


def speaker_documents():
    doc = Document()
    document_style(doc, 'Task-Aware 算法汇报｜逐页讲稿与控时')
    doc.add_heading('10–13 分钟逐页讲稿', 0)
    doc.add_paragraph('状态：作者代码核对工作稿。主讲 12 页，计划 710 秒（11 分 50 秒）；备份 5 页不计入。'
                      '这是建议时长，不是实际录音测量。论文 v2、原文图表和定理假设尚待核对。')
    markdown = ['# 10–13 分钟逐页讲稿', '', '计划 11 分 50 秒；主讲 12 页，备份 5 页。工作稿，待 v2 核对。', '']
    elapsed = 0
    for item in SLIDES:
        if item['seconds']:
            end = elapsed + item['seconds']
            label = f'{elapsed//60:02d}:{elapsed%60:02d}–{end//60:02d}:{end%60:02d}'
            elapsed = end
        else:
            label = '问答备用'
        title = f'{item["id"]}  {item["title"]}'
        doc.add_heading(title, 1)
        doc.add_paragraph('时间：' + label)
        doc.add_paragraph(item['notes'])
        doc.add_paragraph('本页带走：' + item['takeaway'])
        doc.add_paragraph('来源：' + item['source'])
        markdown.extend(['## ' + title, '', '**时间：' + label + '**', '', item['notes'], '',
                         '**带走：** ' + item['takeaway'], '', '**来源：** ' + item['source'], ''])
    assert elapsed == 710
    doc.save(OUT / '逐页讲稿与控时.docx')
    (OUT / '逐页讲稿与控时.md').write_text('\n'.join(markdown).rstrip() + '\n')
    (OUT / '待补图方法论文本.md').write_text(SUPPLEMENT.rstrip() + '\n')
    # Editable text briefing: preserve its Markdown headings as Word sections.
    brief = Document()
    document_style(brief, 'Task-Aware｜待补图方法论文本')
    for block in SUPPLEMENT.split('\n\n'):
        if not block.strip():
            continue
        if block.startswith('# '):
            brief.add_heading(block[2:].strip(), 0)
        elif block.startswith('## '):
            brief.add_heading(block[3:].strip(), 1)
        else:
            brief.add_paragraph(block.replace('**', '').replace('`', '').strip())
    brief.save(OUT / '待补图方法论文本.docx')


def source_records():
    destination = OUT / 'assets/original/author_repository_unlearning.png'
    destination.parent.mkdir(parents=True, exist_ok=True)
    raw = SOURCE / 'paper_original_image.png'
    if raw.exists():
        shutil.copyfile(raw, destination)
    assert destination.exists()
    original_hash = sha(destination)
    sources = []
    for filename in ['author_readme.txt', 'author_code/eval_unchange.py',
                     'author_code/eval_unlearn.py', 'author_code/utils/optimization.py',
                     'author_code/utils/net.py', 'author_code/func_operation.py']:
        path = SOURCE / filename
        if path.exists():
            repo_file = 'readme.md' if filename == 'author_readme.txt' else filename.removeprefix('author_code/')
            sources.append(dict(url=AUTHOR.replace('/tree/', '/blob/') + '/' + repo_file,
                                sha256=sha(path), role='author_source_read', accessed='2026-10-22'))
    course = SOURCE / 'mit-ocw'
    for file in ['content/pages/how-to-speak.md', 'content/resources/how-to-speak-1.md', 'data/course.json']:
        path = course / file
        if path.exists():
            sources.append(dict(url='https://github.com/mitocwcontent/res.tll-005-january-iap-2018/blob/'
                                'b46fd4081b4a89d18d6efafef532e7a48e5a0d9e/' + file,
                                sha256=sha(path), role='MIT_official_course_source_read', accessed='2026-10-22'))
    blocked = []
    for file in ['initial_access.json', 'more_access.json']:
        path = SOURCE / file
        if path.exists():
            blocked.extend({k: x[k] for k in ['key', 'url', 'error'] if k in x}
                           for x in json.loads(path.read_text()) if 'error' in x)
    record = dict(status='WORK_DRAFT_PENDING_ARXIV_V2_AND_NSE_PRIMARY_PAGE',
                  date='2026-10-22', timezone='Asia/Shanghai', target_paper=PAPER,
                  requested_duration_minutes=[10, 13], planned_main_seconds=710,
                  paper_full_text_read=False, paper_figure_numbers_verified=False,
                  original_paper_experimental_figures_available=False,
                  source_code_read=True, mit_official_course_page_source_read=True,
                  mit_full_video_watched=False, mit_nse_page_read=False,
                  new_experiment_runs=False, image_generation_called=False,
                  original_author_image=dict(file=str(destination.relative_to(OUT)), sha256=original_hash,
                                             url=AUTHOR.replace('/tree/', '/blob/') + '/repo_src/unlearning.png',
                                             unchanged=True, paper_figure_identity='not_verified'),
                  sources=sources, blocked_access=blocked)
    target = OUT / 'source_manifest.json'
    if sources:
        target.write_text(json.dumps(record, ensure_ascii=False, indent=2) + '\n')
    elif not target.exists():
        raise RuntimeError('Source cache or committed provenance required')
    figure_manifest = dict(figures=[dict(id='A01', file=str(destination.relative_to(OUT)),
                                        type='unchanged_author_repository_original_image', sha256=original_hash,
                                        used_on='S02', original_paper_figure_number=None)],
                           pending=[dict(id='R1', status='requires_user_uploaded_v2_PDF'),
                                    dict(id='R2', status='requires_user_uploaded_v2_PDF')],
                           proposed_explanatory_figures=['D1', 'D2', 'D3'],
                           generated_explanatory_figures=[])
    (OUT / 'figure_manifest.json').write_text(json.dumps(figure_manifest, ensure_ascii=False, indent=2) + '\n')


def design_readme():
    (OUT / 'MIT指南设计核对.md').write_text(f'''# 两份 MIT 指南的设计适配与核验状态

## 来源范围

- [How to Speak]({MIT_SPEAK})：网站直连 403；已通过 MIT 官方课程仓库读取同一课程页面源码。页面明确包括开场、四种启发、工具、承诺 / 激励 / 思考、说服以及结尾等部分。没有取得完整字幕或观看完整视频。
- [NSE CommKit Slideshow]({MIT_SLIDES})：页面直连及 WordPress 接口均 403，尚未取得原页正文。因此以下视觉原则是演示设计归纳，**不冒充该页面逐字摘录，也不声称已经完成两页全文逐项审查。**
- 这是设计核对表，不是 MIT 官方评分、认证或官方两项量化指标。

## A：口头叙事与控时

| 设计要求 | 在本稿中的落点 | 核验方式 |
|---|---|---|
| 开场告诉听众能带走什么 | S01：如何删、如何补偿、Cost 梯度 | 课程页承诺主题可核对；本稿为中文适配 |
| 主线清楚、定义逐步出现 | S03直接遗忘→S04–05 PAMU→S06–10 TAMU | 结构检查 |
| 关键思想重复但层层推进 | S01、S06、S09、S12 | 演讲设计归纳，不引作原页原句 |
| 主讲与答疑推导分开 | 12 主讲页 + 5 备份页 | 页数与时间表检查 |
| 用实质性总结收束 | S12保留三条结论，不以“谢谢”占一页 | 课程页结尾主题可核对；具体做法为适配 |
| 控制 10–13 分钟 | 计划 710 秒，即11分50秒 | 只是计划；仍需实际排练 |

## B：幻灯片的信息组织与可读性

| 设计要求 | 在本稿中的落点 | 当前证据性质 |
|---|---|---|
| 一页一个主张，标题写结论 | S03–S12标题均表达算法主张 | 归纳适配；NSE原页未读 |
| 公式后面有对象、符号与解释 | 每页至多3个主公式；B3约束备用例外 | 程序和渲染检查 |
| 正文以大字呈现，细节移入讲稿 | 主体约21–25pt；来源脚注另计 | 版面检查，不等于官方阈值 |
| 颜色只承担固定语义 | 深蓝文字、青色主线、橙色边界与待核对 | 归纳适配 |
| 原始图保留出处 | S02使用作者图片，未核对v2图号 | 来源哈希检查 |
| 用空白和对齐建立层次 | 16:9、固定标题 / 带走句 / 来源栏 | PDF视觉检查 |
| 缺失证据不靠装饰或生成结果填充 | R1/R2保留待PDF状态；D1–D3只给方法文本 | 科学溯源检查 |

## 正式定稿前还需做什么

取得 v2 PDF 核对算法、定理与原图；取得 NSE 页正文后再做逐条标准核对。补齐结果图可能需要用一张结果页替换当前 S11，重新分配主讲时间，避免直接增加时长。
'''.rstrip() + '\n')
    (OUT / '原文核对清单.md').write_text('''# 原文 v2 定稿核对清单

目前已完成的是可审阅的作者代码核对工作稿，未完成原文精读终稿。

- [ ] 取得 arXiv:2308.14412v2 PDF，记录文件 SHA256、页数、标题和版本日期。
- [ ] 核对直接遗忘的具体损失归一化、Hessian取值及公式编号。
- [ ] 核对 PAMU 的原文目标、权重变量含义、线性近似与约束集合。
- [ ] 核对 TAMU 三层问题具体写法；当前稿只用讲解重述，无原文编号。
- [ ] 核对梯度存在性定理的假设、结论范围及不光滑边界。
- [ ] 核对两阶段成本和约束是否与代码一致，代码修复不等于论文原式。
- [ ] 列出原文全部图表，选取支持主讲主张的最少图，填写图号、页码、图注与结论。
- [ ] 确认作者仓库 unlearning.png 是否与v2某图相同；未确认前只称作者仓库原图。
- [ ] 核对原文实验划分、模型、遗忘比例、重加权预算、基线及量化结果。
- [ ] 移除工作稿提示，或保留仍未完成的明确边界；不能在缺失时直接改成“定稿”。

本稿不使用此前独立复现的MSE/MAPE/Cost曲线充当原论文图；也没有将输入温度扰动实验写成原文实验。
'''.rstrip() + '\n')
    (OUT / 'README.md').write_text(f'''# Task-Aware 论文算法汇报｜10–13 分钟工作稿

**当前是可审阅工作稿，尚不是原文核对终稿。** 用户指定的 v2 论文与 NSE 指南页面被网络代理403拦截。

[下载完整包]({BASE}/exports/task_aware_paper_talk.zip?raw=true) ·
[下载PPT]({STEM}.pptx?raw=true) · [阅读PDF]({STEM}.pdf)

## 已完成

- 12张主讲页 + 5张备份页，计划11分50秒（不是实测语速）。
- 重点：直接遗忘、PAMU、权重LP、TAMU、两阶段Cost、KKT隐式求导、iHVP样本评分、最终带权MSE更新。
- 每页备注内附讲稿、时长、公式源码与来源。
- 作者仓库原始流程图原样插入并标注来源，未冒称已确认的v2图号。
- [待补图方法文本](待补图方法论文本.md)，含3份可交给制图工具的说明；没有调用图像生成。
- [讲稿与控时](逐页讲稿与控时.md)、[MIT设计核对](MIT指南设计核对.md)、[原文核对清单](原文核对清单.md)。

## 仍然缺少的材料

**请上传 arXiv:2308.14412v2 PDF。** 正式稿需要核对原文公式、梯度定理和结果图。当前没有原文实验图，只有作者仓库原图。

How to Speak 已取得 MIT 官方课程页源码，未观看全视频；NSE Slideshow 正文未取得，视觉设计为归纳适配，不声称逐条验证其原页标准。

## 文件

- `{STEM}.pptx` / `.pdf`：演示工作稿。
- `逐页讲稿与控时.docx` / `.pdf` / `.md`：独立讲稿。
- `待补图方法论文本.docx` / `.pdf` / `.md`：可选示意图需求。
- `slides.json` / `equations.json`：可重复生成的内容与公式。
- `figure_manifest.json` / `source_manifest.json`：原图和访问证据。
- `validation.json`：页数、讲稿、原图一致性与渲染检查。
- `SHA256SUMS`：文件校验。

重新生成：`python reports/build_task_aware_talk.py`。
依赖：python-pptx、python-docx、matplotlib、Pillow、PyMuPDF、soffice、Noto CJK字体。

分支：`{BRANCH}`。没有修改原训练/遗忘代码，没有重新训练或计算Cost，没有合并master。
'''.rstrip() + '\n')


def render():
    targets = [OUT / (STEM + '.pptx'), OUT / '逐页讲稿与控时.docx', OUT / '待补图方法论文本.docx']
    for target in targets:
        result = subprocess.run(['soffice', '-env:UserInstallation=file:///tmp/task-aware-talk-lo',
                                 '--headless', '--convert-to', 'pdf', '--outdir', str(OUT), str(target)],
                                capture_output=True, text=True, timeout=110,
                                env={**os.environ, 'XDG_CACHE_HOME': '/workspace/.cache'})
        print(result.stdout.strip(), flush=True)
        if result.returncode or not target.with_suffix('.pdf').exists():
            raise RuntimeError(result.stdout + '\n' + result.stderr)


def normalize(text):
    value = unicodedata.normalize('NFKC', text)
    return re.sub(r'[\s“”‘’]', '', value).replace('−', '-').replace('–', '-')


def inspect():
    prs = Presentation(OUT / (STEM + '.pptx'))
    pdf = fitz.open(OUT / (STEM + '.pdf'))
    assert len(prs.slides) == len(pdf) == len(SLIDES) == 17
    assert len([x for x in SLIDES if x['seconds']]) == 12
    assert sum(x['seconds'] for x in SLIDES) == 710
    original = OUT / 'assets/original/author_repository_unlearning.png'
    original_matches = []
    text_missing = []
    out_of_page = []
    page_images = []
    for i, (slide, page, item) in enumerate(zip(prs.slides, pdf, SLIDES)):
        assert item['notes'] in slide.notes_slide.notes_text_frame.text
        rendered_text = normalize(page.get_text())
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width + 1000
            assert shape.top + shape.height <= prs.slide_height + 1000
            if shape.has_text_frame:
                chunk = normalize(shape.text)
                if chunk and chunk not in rendered_text:
                    text_missing.append(dict(slide=i + 1, text=shape.text))
            if shape.shape_type == 13 and shape.image.blob == original.read_bytes():
                original_matches.append(i + 1)
        for block in page.get_text('dict')['blocks']:
            if block['type'] != 0:
                continue
            for line in block['lines']:
                for span in line['spans']:
                    r = fitz.Rect(span['bbox'])
                    if r.x0 < -.5 or r.y0 < -.5 or r.x1 > page.rect.width + .5 or r.y1 > page.rect.height + .5:
                        out_of_page.append(dict(page=i + 1, text=span['text'], bbox=list(r)))
        path = QA / f'slide_{i+1:02d}.png'
        page.get_pixmap(matrix=fitz.Matrix(1.1, 1.1), alpha=False).save(path)
        page_images.append(path)
    assert original_matches == [2]
    assert not text_missing, text_missing
    assert not out_of_page, out_of_page
    # A compact contact sheet for manual inspection; no new scientific figure.
    w, h, cols = 400, 250, 3
    sheet = Image.new('RGB', (w * cols, h * ((len(page_images) + cols - 1) // cols)), '#DEE5E9')
    draw = ImageDraw.Draw(sheet)
    for i, p in enumerate(page_images):
        with Image.open(p) as im:
            thumb = ImageOps.contain(im, (w - 8, h - 24))
            x, y = i % cols * w + (w - thumb.width) // 2, i // cols * h + 3
            sheet.paste(thumb, (x, y))
            draw.text((i % cols * w + 9, i // cols * h + h - 18), SLIDES[i]['id'], fill='black')
    sheet.save(QA / 'slides_contact.png')
    documents = []
    for path in [OUT / '逐页讲稿与控时.pdf', OUT / '待补图方法论文本.pdf']:
        d = fitz.open(path)
        assert all(len(p.get_text().strip()) > 30 for p in d)
        documents.append(dict(file=path.name, pages=len(d), sha256=sha(path)))
        d.close()
    result = dict(status='WORK_DRAFT_NOT_PRIMARY_PAPER_VERIFIED', slides_total=17, main_slides=12,
                  backup_slides=5, planned_seconds=710, actual_spoken_duration_measured=False,
                  notes_present_on_all_slides=True, original_author_image_used_on_slides=original_matches,
                  original_image_byte_identical=True, text_missing=text_missing, out_of_page_text=out_of_page,
                  original_paper_result_figures_present=False, arxiv_v2_full_text_read=False,
                  generated_supplementary_figures=False, helper_documents=documents,
                  contact_sheet=str(QA / 'slides_contact.png'),
                  powerpoint_sha256=sha(OUT / (STEM + '.pptx')), pdf_sha256=sha(OUT / (STEM + '.pdf')))
    (OUT / 'validation.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps(result, ensure_ascii=False, indent=2), flush=True)


def package():
    files = sorted(p for p in OUT.rglob('*') if p.is_file() and p.name != 'SHA256SUMS')
    (OUT / 'SHA256SUMS').write_text(''.join(sha(p) + '  ' + p.relative_to(OUT).as_posix() + '\n' for p in files))
    target = ROOT / 'exports/task_aware_paper_talk.zip'
    with ZipFile(target, 'w', ZIP_DEFLATED) as z:
        for p in sorted(OUT.rglob('*')):
            if p.is_file():
                z.write(p, 'task_aware_paper_talk/' + p.relative_to(OUT).as_posix())
        for name in ['reports/build_task_aware_talk.py', 'reports/task_aware_talk_content.py']:
            z.write(ROOT / name, 'repository_scripts/' + name)
    with ZipFile(target) as z:
        assert z.testzip() is None
        for p in OUT.rglob('*'):
            if p.is_file():
                assert z.read('task_aware_paper_talk/' + p.relative_to(OUT).as_posix()) == p.read_bytes()
    target.with_suffix('.zip.sha256').write_text(sha(target) + '  ' + target.name + '\n')
    print(dict(zip=str(target), bytes=target.stat().st_size, sha256=sha(target)), flush=True)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    QA.mkdir(parents=True, exist_ok=True)
    source_records()
    formula_assets()
    build_presentation()
    speaker_documents()
    design_readme()
    render()
    inspect()
    package()


if __name__ == '__main__':
    main()
