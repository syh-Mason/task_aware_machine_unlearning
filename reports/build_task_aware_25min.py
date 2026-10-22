"""Build paper-verified 25-minute presentation and editable speaker notes."""
from copy import deepcopy
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import unicodedata
from zipfile import ZipFile, ZIP_DEFLATED

os.environ.setdefault('MPLCONFIGDIR', '/workspace/.cloud-setup/task-aware/matplotlib')
os.environ.setdefault('XDG_CACHE_HOME', '/workspace/.cache')

import fitz
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from PIL import Image, ImageDraw, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches as DInches, Pt as DPt

from task_aware_25min_content import SLIDES, OPTIONAL_FIGURE_BRIEF

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'reports/task_aware_25min'
SOURCE = Path('/workspace/.cloud-setup/task-aware/talk25-sources')
QA = Path('/workspace/.cloud-setup/task-aware/talk25-qa')
STEM = 'TaskAware_25分钟_论文与特征分析'
FONT = 'Noto Sans CJK SC'
NAVY, TEAL, ORANGE = '16354A', '087F8C', 'B75331'
TEXT, GRAY, LIGHT, PALE = '243D4E', '596B77', 'EAF3F6', 'FBF2EC'
BRANCH = 'slides/task-aware-25min-paper'
BASE = f'https://github.com/syh-Mason/task_aware_machine_unlearning/blob/{BRANCH}'


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def color(hexcolor):
    return RGBColor.from_string(hexcolor)


def rect(slide, x, y, w, h, fill, line=None, rounded=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if line:
        shape.line.color.rgb = color(line)
        shape.line.width = Pt(.7)
    else:
        shape.line.fill.background()
    if rounded:
        shape.adjustments[0] = .05
    return shape


def text(slide, value, x, y, w, h, size=22, bold=False, ink=TEXT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    for i, line in enumerate(value.split('\n')):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.space_after = Pt(5)
        p.space_before = Pt(0)
        p.line_spacing = 1.10
        r = p.add_run()
        r.text = line
        r.font.name, r.font.size = FONT, Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color(ink)
        pr = r._r.get_or_add_rPr()
        ea = OxmlElement('a:ea')
        ea.set('typeface', FONT)
        pr.append(ea)
    return box


def picture(slide, path, x, y, w, h):
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    width, height = iw * scale, ih * scale
    shape = slide.shapes.add_picture(str(path), Inches(x + (w - width) / 2),
                                     Inches(y + (h - height) / 2),
                                     width=Inches(width), height=Inches(height))
    return shape


def formula_assets():
    folder = OUT / 'assets/equations'
    folder.mkdir(parents=True, exist_ok=True)
    manifest = []
    for slide in SLIDES:
        for i, expression in enumerate(slide.get('equations', [])):
            expression = expression.replace(r'\bigg|_{0}', r'|_{0}').replace(r'\mathbf1', r'\mathbf{1}').replace(r'\xrightarrow{\ (12)\ }', r'\rightarrow').replace(r'\ne0', r'\neq0')
            expression = re.sub(r'\\le(?![a-zA-Z])', r'\\leq', expression)
            expression = re.sub(r'\\ge(?![a-zA-Z])', r'\\geq', expression)
            stem = f'{slide["id"]}_{i+1:02d}'
            # Typeset directly to the text's extent, without editing source art.
            fig = plt.figure(figsize=(18, 1.8), facecolor='white')
            item = fig.text(0, .5, expression, fontsize=34, va='center', color='#' + NAVY)
            fig.canvas.draw()
            extent = item.get_window_extent(fig.canvas.get_renderer()).transformed(fig.dpi_scale_trans.inverted())
            extent.x0 -= .05
            extent.x1 += .05
            extent.y0 -= .05
            extent.y1 += .05
            for ext in ['png', 'svg']:
                fig.savefig(folder / (stem + '.' + ext), dpi=240, bbox_inches=extent, facecolor='white')
                if ext == 'svg':
                    path = folder / (stem + '.svg')
                    path.write_text('\n'.join(line.rstrip() for line in path.read_text().splitlines()) + '\n')
            plt.close(fig)
            manifest.append(dict(slide=slide['id'], equation_index=i + 1, latex=expression,
                                 provenance='Explanatory math restatement, not an extracted paper equation or figure',
                                 png='assets/equations/' + stem + '.png'))
    (OUT / 'equations.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n')


def header(slide, item, number):
    rect(slide, 0, 0, 13.333, .10, TEAL if item['seconds'] else GRAY)
    text(slide, 'TASK-AWARE  /  ' + item['section'], .55, .22, 8.8, .28, 11, True, TEAL)
    timing = f'{item["seconds"]} 秒' if item['seconds'] else '问答备用 · 不计主讲'
    text(slide, timing, 10.4, .22, 2.35, .28, 11, False, GRAY)
    text(slide, item['title'], .55, .72, 12.25, .70, 26, True, NAVY)
    rect(slide, .55, 6.37, 12.23, .57, LIGHT)
    text(slide, item['takeaway'], .76, 6.47, 11.80, .42, 16, True, TEAL)
    text(slide, item['source'], .55, 7.10, 11.82, .26, 8, False, GRAY)
    text(slide, f'{number:02d}', 12.47, 7.08, .4, .28, 10, False, GRAY)


def add_equations(slide, item, y=1.58, gap=.75, maxheight=.57):
    for i, _ in enumerate(item.get('equations', [])):
        path = OUT / 'assets/equations' / f'{item["id"]}_{i+1:02d}.png'
        shape = picture(slide, path, .85, y + i * gap, 11.60, maxheight)
        shape.name = f'{item["id"]}:equation:{i+1}'
        shape._element.nvPicPr.cNvPr.set('descr', item['equations'][i])


def add_lines(slide, lines, y, gap=.60, size=22, x=.85, w=11.55):
    for i, line in enumerate(lines):
        text(slide, line, x, y + i * gap, w, .53, size, False, TEXT)



PDF_SOURCE = Path('/workspace/attachments/aef60b13-9ff5-4c3a-9cd2-3c2f9b8a9bcc/Task-Aware_Machine_Unlearning_and_Its_Application_in_Load_Forecasting.pdf')
PDF_HASH = '24baa12c60fc48b88f6fc83988183d12a2991cb846b22179840f17fcf9568627'
CROPS = {
 'fig1': (3, (70,63,265,178)), 'fig2': (5,(305,61,556,139)),
 'fig3': (6,(38,62,290,132)), 'fig4': (7,(42,62,557,188)),
 'fig5': (7,(42,237,557,371)), 'fig6': (8,(37,63,553,177)),
 'fig7': (8,(301,220,553,332)), 'fig8': (11,(42,63,557,288)),
 'fig9': (11,(42,338,557,565)), 'table2': (10,(300,62,553,183)),
 'table3': (10,(300,192,553,324)),
 'fig8_cost_top': (11,(391,63,557,175)),
 'fig8_cost_bottom': (11,(391,177,557,288)),
}

def extract():
    assert sha(PDF_SOURCE) == PDF_HASH
    folder = OUT / 'assets/original'; folder.mkdir(parents=True, exist_ok=True)
    source = fitz.open(PDF_SOURCE)
    records = []
    for name, (n, bbox) in CROPS.items():
        clip = fitz.Rect(bbox)
        d = fitz.open(); page = d.new_page(width=clip.width, height=clip.height)
        page.show_pdf_page(page.rect, source, n-1, clip=clip)
        page.get_pixmap(matrix=fitz.Matrix(5,5), alpha=False).save(folder/(name+'.png'))
        # Extracted figure vectors, with the source PDF kept private in the workspace.
        svg = page.get_svg_image(text_as_path=True)
        (folder/(name+'.svg')).write_text('\n'.join(s.rstrip() for s in svg.splitlines())+'\n')
        records.append(dict(asset=name, source_pdf_sha256=PDF_HASH, pdf_page=n,
                            journal_page=7177+n, clip_points=list(bbox),
                            operation='rectangular crop; no curve or number edited',
                            used_on=[s['id'] for s in SLIDES if s.get('figure')==name or (s['kind']=='nn_panels' and name.startswith('fig8_cost'))]))
        d.close()
    (OUT/'figure_manifest.json').write_text(json.dumps(records,ensure_ascii=False,indent=2)+'\n')

def art(slide, name, x,y,w,h):
    pic = picture(slide, OUT/'assets/original'/f'{name}.png', x,y,w,h)
    pic.name = 'Paper original: '+name
    pic._element.nvPicPr.cNvPr.set('descr', 'Xu & Teng 2024, '+name+'; cropped from uploaded final IEEE PDF; see figure_manifest.json')

def lines(slide, values, x=.85,y=3.8,w=11.6,gap=.62,size=21,h=.6):
    for i, value in enumerate(values):
        text(slide, value, x,y+i*gap,w,h,size)

def build():
    prs=Presentation(); prs.slide_width=Inches(13.333333); prs.slide_height=Inches(7.5)
    for number, item in enumerate(SLIDES,1):
        slide=prs.slides.add_slide(prs.slide_layouts[6]); kind=item['kind']
        if kind=='cover':
            rect(slide,0,0,.16,7.5,TEAL)
            text(slide,'论文精读 / 算法推导 / 研究扩展',.75,.65,11.8,.4,16,True,TEAL)
            text(slide,item['title'],.75,1.5,11.8,1.55,36,True,NAVY)
            text(slide,item['subtitle'],.78,3.2,11.7,.75,21,False,GRAY)
            text(slide,item['lines'][0],.78,4.25,11.7,.6,18,False,GRAY)
            text(slide,'直接遗忘 → 任务补偿 → 原文证据 → 特征分析的可能性',.78,5.12,11.7,.8,23,True,TEAL)
            text(slide,'25 分钟计划 · 27 页主讲 + 7 页备份 · 依据上传的正式发表版',.78,6.48,11.7,.4,16,False,GRAY)
            text(slide,'W. Xu & F. Teng, IEEE TPWRS, 2024 · DOI: 10.1109/TPWRS.2024.3376828',.78,7.06,11.7,.3,11,False,GRAY)
        else:
            header(slide,item,number)
            if kind in ['equations','comparison']:
                n=len(item.get('equations',[])); gap=.73 if n>2 else .85
                add_equations(slide,item,y=1.62,gap=gap,maxheight=.55 if n>2 else .64)
                start=1.62+n*gap+.15
                lines(slide,item['lines'],y=start,gap=.59,size=20 if n>=3 or kind=='comparison' else 21,h=.57)
            elif kind=='figure_side':
                if item['figure']=='fig3':
                    art(slide,'fig3',.8,1.55,11.65,2.35)
                    lines(slide,item['lines'],y=4.08,gap=.60,size=21)
                else:
                    art(slide,item['figure'],.65,1.55,7.5,4.55)
                    rect(slide,8.45,1.68,4.28,4.35,LIGHT,rounded=True)
                    lines(slide,item['lines'],x=8.7,y=1.98,w=3.8,gap=1.25,size=20,h=1.14)
            elif kind=='figure_equations':
                art(slide,item['figure'],.8,1.53,11.7,1.9)
                add_equations(slide,item,y=3.55,gap=.78,maxheight=.68)
                lines(slide,item['lines'],y=4.40,gap=.58,size=20)
            elif kind=='figure_wide':
                art(slide,item['figure'],.65,1.56,12.0,3.25)
                lines(slide,item['lines'],y=5.00,gap=.54,size=19,h=.66)
            elif kind=='figure_full':
                art(slide,item['figure'],.7,1.50,11.9,4.32)
                lines(slide,item['lines'],y=5.90,gap=.5,size=17,h=.38)
            elif kind=='nn_panels':
                art(slide,'fig8_cost_top',.7,1.70,5.8,3.75)
                art(slide,'fig8_cost_bottom',6.85,1.70,5.8,3.75)
                text(slide,'完整遗忘：蓝剩余 / 红删除 / 绿测试；虚线原始',.80,5.56,5.68,.65,16,False,GRAY)
                text(slide,'补偿：红 MSE-PAMU / 蓝 MAPE-PAMU / 绿 TAMU',6.85,5.56,5.68,.65,16,False,GRAY)
            elif kind=='procedure':
                lines(slide,item['lines'],y=1.62,gap=.72,size=21,h=.64)
                add_equations(slide,item,y=5.44,maxheight=.52)
            elif kind in ['protocol','audit','sources']:
                lines(slide,item['lines'],y=1.75,gap=1.02,size=22,h=.88)
            elif kind=='closing':
                for i,line in enumerate(item['lines']):
                    text(slide,f'0{i+1}',.85,1.85+i*1.35,.8,.65,30,True,TEAL)
                    text(slide,line,1.95,1.90+i*1.35,10.35,1.0,24,True,NAVY)
            elif kind=='tables':
                art(slide,'table2',.75,1.60,5.7,3.9)
                art(slide,'table3',6.85,1.60,5.7,3.9)
                lines(slide,item['lines'],y=5.58,gap=.36,size=16,h=.36)
            else: raise ValueError(kind)
        slide.notes_slide.notes_text_frame.text=(f'{item["id"]} | 计划 {item["seconds"]} 秒\n\n'+item['notes']+'\n\n来源：'+item['source']+'\n\n公式源文本：\n'+'\n'.join(item.get('equations',[])))
    prs.core_properties.title='Task-Aware Machine Unlearning｜25分钟算法与特征分析'
    prs.core_properties.subject='IEEE正式版原图、算法解析、剩余集重训比较及特征研究扩展'
    prs.core_properties.author='Task-Aware 文献汇报'
    prs.save(OUT/(STEM+'.pptx'))
    (OUT/'slides.json').write_text(json.dumps(SLIDES,ensure_ascii=False,indent=2)+'\n')

def docs():
    doc=Document()
    normal=doc.styles['Normal']; normal.font.name=FONT; normal.font.size=DPt(10.5)
    normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
    for k in ['Title','Heading 1','Heading 2']:
        doc.styles[k].font.name=FONT
        doc.styles[k]._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
    doc.add_heading('Task-Aware｜25分钟逐页讲稿',0)
    doc.add_paragraph('27页主讲计划1500秒，另有7页问答备份。时间为演讲计划，未做实际录音测量。主依据为上传的12页IEEE正式版；扩展研究与原文结果分别标注。')
    md=['# 25分钟逐页讲稿','', '27页主讲 + 7页备份；计划25:00，并非录音测量。','']
    elapsed=0
    for i,s in enumerate(SLIDES,1):
        end=elapsed+s['seconds']; time=f'{elapsed//60:02d}:{elapsed%60:02d}–{end//60:02d}:{end%60:02d}' if s['seconds'] else '问答备份'
        if s['seconds']: elapsed=end
        title=f'{i:02d} / {s["id"]} {s["title"]}'
        doc.add_heading(title,1); doc.add_paragraph('建议时间：'+time); doc.add_paragraph(s['notes'])
        doc.add_paragraph('带走：'+s['takeaway']); doc.add_paragraph('来源：'+s['source'])
        md += ['## '+title,'',time,'',s['notes'],'','带走：'+s['takeaway'],'','来源：'+s['source'],'']
    assert elapsed==1500
    doc.save(OUT/'逐页讲稿与控时.docx'); (OUT/'逐页讲稿与控时.md').write_text('\n'.join(md).rstrip()+'\n')
    brief=OPTIONAL_FIGURE_BRIEF+'''\n## 新增可选图：从样本影响走向特征分析\n\n用途：X02–X03，属于研究方案，不是论文实证。\n\n生成方法论文本：绘制16:9白底科研示意图。上排表示训练链：训练样本(x_i,y_i)→按样本加权的MSE→参数θ。下排表示推理与任务链：测试输入x的温度等10类特征→冻结特征提取器h→线性输出头→14节点预测负荷→预测驱动调度→固定真实负荷的补救→Cost。用三种不同颜色标记：样本权重影响、测试输入敏感性、隐藏层归因。训练记录中的特征扰动用虚线经参数θ接到任务链，标明需要混合二阶导数与局部隐式求导。不要画成训练权重直接乘到测试输入；不要声称特征归因等于特征遗忘或因果发现；不添加数值或虚构曲线。底注“扩展研究方案，非原论文结果”。\n\n## 可继续阅读的文献和工具\n\n- Koh & Liang, 2017, Understanding Black-box Predictions via Influence Functions：https://arxiv.org/abs/1703.04730 （样本影响的基础文献）\n- Sundararajan et al., 2017, Axiomatic Attribution for Deep Networks：https://arxiv.org/abs/1703.01365 （输入归因；扩展阅读，非原文算法）\n- Captum 官方：https://captum.ai/ （输入与层归因工具；不等于机器遗忘保证）\n- 原论文DOI：https://doi.org/10.1109/TPWRS.2024.3376828\n- MIT How to Speak：https://ocw.mit.edu/courses/res-tll-005-how-to-speak-january-iap-2018/pages/how-to-speak/\n- MIT NSE Slideshow：https://mitcommlab.mit.edu/nse/commkit/slideshow/\n\n这些链接用于继续阅读，不表示本轮已重新访问其全部原文。\n'''
    (OUT/'可选补图方法论与阅读链接.md').write_text(brief)
    d=Document(); d.styles['Normal'].font.name=FONT; d.styles['Normal']._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
    for block in brief.split('\n\n'):
        if block.startswith('# '): d.add_heading(block[2:],0)
        elif block.startswith('## '): d.add_heading(block[3:],1)
        else: d.add_paragraph(block.replace('**','').replace('`',''))
    d.save(OUT/'可选补图方法论与阅读链接.docx')
    readme=f'''# Task-Aware：25分钟论文汇报\n\n依据：上传IEEE最终发表版，12页，DOI 10.1109/TPWRS.2024.3376828。\n\n- `{STEM}.pptx`：可编辑演示文稿；27页主讲+7页备份，计划25分钟。\n- `{STEM}.pdf`：投影/预览版。\n- `逐页讲稿与控时.docx`：逐页说明与时间建议，PPT备注亦含讲稿。\n- `可选补图方法论与阅读链接.docx`：可交给生成工具的补图描述；本轮未生成实验图。\n- `figure_manifest.json`：原论文图1–9及结构表的裁剪来源与坐标。\n- `equations.json`、`slides.json`：公式源码及内容，支持核验。\n\n## 本次新增重点\n\n1. 剩余集均匀重训与TAMU加权拟合的目标差异；直接MU的速度差异与精确性条件。\n2. 从样本影响到输入敏感性、隐藏层归因、训练特征影响的三种路径；不将归因冒充特征遗忘。\n3. 样本/特征/Cost的可检验扩展方案，明确属于研究建议，非论文实证。\n\n## 模板参考限制\n\nB站视频BV1WF411Y73C页面及接口均因代理403无法访问，未观看、未复刻其模板。暂用白底深蓝/青色学术版式，沿用此前MIT开场承诺、每页一个结论、算法递进与实质总结原则。MIT NSE页面此前不可访问，不声称逐项核验。提供视频关键页截图后可进一步匹配视觉模板。\n\n## 科学口径\n\n原图与实验数字来自上传的IEEE版，未重新运行训练或Cost实验。原文式18不含首阶段可行性松弛惩罚，不能无条件等同此前代码的两个求解目标之和。原文时间划分为前80%训练/后20%测试；图6删除25%。冻结特征+满秩线性MSE的精确性不推广到全CNN；参数距离不等于认证遗忘。KKT导数的活跃集边界与原文21/22的符号一致性在备注和备份中说明。原论文整份PDF未重新公开，只提取署名引用的汇报所需图表。\n\n## 重新生成\n\n安装python-pptx、python-docx、pymupdf、matplotlib、pillow；系统需LibreOffice及Noto中文字体。运行`python reports/build_task_aware_25min.py`，需本地原论文PDF（路径在脚本PDF_SOURCE中，可自行修改）；环境信息见build_environment.json。\n'''
    (OUT/'README.md').write_text(readme)

def render():
    for name in [STEM+'.pptx','逐页讲稿与控时.docx','可选补图方法论与阅读链接.docx']:
        r=subprocess.run(['soffice','-env:UserInstallation=file:///tmp/task-aware-25-lo','--headless','--convert-to','pdf','--outdir',str(OUT),str(OUT/name)],capture_output=True,text=True,timeout=150,env={**os.environ,'XDG_CACHE_HOME':'/workspace/.cache'})
        print(r.stdout.strip(),flush=True)
        if r.returncode or not (OUT/Path(name).with_suffix('.pdf')).exists(): raise RuntimeError(r.stdout+r.stderr)

def normalize(t):
    return re.sub(r'[\s“”‘’]', '',unicodedata.normalize('NFKC',t)).replace('−','-').replace('–','-')

def inspect():
    QA.mkdir(parents=True,exist_ok=True)
    prs=Presentation(OUT/(STEM+'.pptx')); pdf=fitz.open(OUT/(STEM+'.pdf'))
    assert len(prs.slides)==len(pdf)==34
    missing=[]; badbounds=[]; pics=[]; original_checks=[]
    for i,(slide,page,s) in enumerate(zip(prs.slides,pdf,SLIDES)):
        assert s['notes'] in slide.notes_slide.notes_text_frame.text
        rendered=normalize(page.get_text())
        for shape in slide.shapes:
            if shape.left<0 or shape.top<0 or shape.left+shape.width>prs.slide_width+1000 or shape.top+shape.height>prs.slide_height+1000: badbounds.append([i+1,shape.name])
            if shape.has_text_frame and normalize(shape.text) and normalize(shape.text) not in rendered:
                missing.append(dict(slide=i+1,text=shape.text))
            if shape.shape_type == 13 and shape.name.startswith('Paper original: '):
                name=shape.name.split(': ')[1]
                assert shape.image.blob == (OUT/'assets/original'/(name+'.png')).read_bytes()
                original_checks.append(dict(slide=i+1, asset=name, exact_embedded_png_match=True))
        path=QA/f'slide_{i+1:02d}.png'; page.get_pixmap(matrix=fitz.Matrix(1,1),alpha=False).save(path); pics.append(path)
    for start in [0,18]:
        subset=pics[start:start+18]; sheet=Image.new('RGB',(1200,250*((len(subset)+2)//3)),'#DEE5E9'); draw=ImageDraw.Draw(sheet)
        for k,path in enumerate(subset):
            im=Image.open(path); thumb=ImageOps.contain(im,(392,220)); x=(k%3)*400+(400-thumb.width)//2; y=(k//3)*250+2
            sheet.paste(thumb,(x,y)); draw.text(((k%3)*400+10,(k//3)*250+227),f'{start+k+1:02d} {SLIDES[start+k]["id"]}',fill='black')
        sheet.save(QA/f'contact_{start+1:02d}.png')
    report=dict(slides=len(pdf),main_slides=27,backup_slides=7,planned_seconds=sum(s['seconds'] for s in SLIDES),missing_text=missing,out_of_page_shapes=badbounds,notes_present=True,paper_sha256=PDF_HASH,video_viewed=False,new_experiments_run=False,original_image_checks=original_checks)
    (OUT/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print('QA',json.dumps(report,ensure_ascii=False),flush=True)
    assert not badbounds,badbounds
    assert not missing,missing

def pack():
    import importlib.metadata
    provenance = dict(
        primary_paper='Xu & Teng, IEEE TPWRS 39(6), 7178–7189, 2024',
        doi='10.1109/TPWRS.2024.3376828',
        uploaded_pdf_sha256=PDF_HASH, full_paper_pages=12,
        requested_main_minutes=25, planned_seconds=1500,
        timing_is_measured=False,
        video='https://www.bilibili.com/video/BV1WF411Y73C',
        video_access='Page and API blocked by proxy; not watched; template not matched',
        editorial_extensions=['uniform retraining versus task-guided weighting',
                              'input sensitivity versus layer attribution versus training-feature influence'],
        new_experiments_run=False)
    (OUT/'source_status.json').write_text(json.dumps(provenance,ensure_ascii=False,indent=2)+'\n')
    (OUT/'build_environment.json').write_text(json.dumps({p:importlib.metadata.version(p) for p in ['python-pptx','python-docx','pymupdf','matplotlib','pillow']},indent=2)+'\n')
    checks={str(p.relative_to(OUT)):sha(p) for p in sorted(OUT.rglob('*')) if p.is_file() and p.name!='checksums.json'}
    (OUT/'checksums.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2)+'\n')
    target=ROOT/'exports/task_aware_25min.zip'; target.parent.mkdir(exist_ok=True)
    with ZipFile(target,'w',ZIP_DEFLATED) as z:
        for p in sorted(OUT.rglob('*')):
            if p.is_file(): z.write(p,str(Path('task_aware_25min')/p.relative_to(OUT)))
        for name in ['build_task_aware_25min.py','task_aware_25min_content.py']:
            z.write(ROOT/'reports'/name,'scripts/'+name)
    target.with_suffix('.zip.sha256').write_text(sha(target)+'  '+target.name+'\n')
    print('ZIP',target,target.stat().st_size,flush=True)

if __name__=='__main__':
    OUT.mkdir(parents=True,exist_ok=True)
    extract(); print('original figures extracted',flush=True)
    formula_assets(); print('equations typeset',flush=True)
    build(); docs(); render(); inspect(); pack()
